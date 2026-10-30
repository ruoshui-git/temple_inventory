"""Report/export contracts exercised against real native stock and permissions."""

# ruff: noqa: RUF001 — Chinese UI punctuation is intentional.

import copy
import csv
import io
import json
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch
from uuid import uuid4

import frappe
from frappe.utils import nowdate
from openpyxl import load_workbook

from temple_inventory import browse, counts, exports, native_api, native_files
from temple_inventory.export_columns import contract as export_contract
from temple_inventory.export_columns import label as export_label
from temple_inventory.projections import Projection, report_page
from temple_inventory.tests.helpers import NativeReportTestCase
from temple_inventory.uom_labels import uom_label


class NativeReportsExportsTests(NativeReportTestCase):
	def test_checked_in_export_column_contract_is_current(self):
		path = Path(__file__).resolve().parents[2] / "frontend/tests/fixtures/export-columns.json"
		self.assertEqual(
			json.loads(path.read_text(encoding="utf-8")),
			export_contract(),
			(
				"Export contract is stale; regenerate with scripts/dx python -c '"
				"import json; from temple_inventory.export_columns import contract; print(json.dumps(contract(), "
				"ensure_ascii=False, indent=2, sort_keys=True))' > frontend/tests/fixtures/export-columns.json"
			),
		)

	def test_loan_history_accepts_outstanding_export_column(self):
		self.receipt(12)
		self.loan(5)
		query = self.query("loans", "items")
		columns = [
			"item_code",
			"item_name",
			"document",
			"kind",
			"date",
			"activity_title",
			"borrower",
			"docstatus",
			"source_warehouse",
			"destination_warehouse",
			"quantity",
			"stock_uom",
			"outstanding",
		]
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(
				query,
				"loan_history",
				columns=columns,
			)
		request = frappe.get_doc("Inventory Export Request", result["name"])
		self.assertEqual(json.loads(request.columns_json), columns)

	def test_default_browse_export_columns_are_accepted(self):
		cases = [
			(
				self.query("inventory"),
				"summary",
				["item_name", "item_code", "available", "stock_uom"],
			),
			(
				self.query("inventory"),
				"locations",
				[
					"item_name",
					"item_code",
					"available",
					"stock_uom",
					"category",
					"warehouse_label",
					"batch_no",
					"expiry_date",
					"remaining_days",
					"expiry_state",
				],
			),
			(
				self.query("expiry", "batches"),
				"batches",
				[
					"item_name",
					"item_code",
					"category",
					"batch_no",
					"expiry_date",
					"remaining_days",
					"stock_qty",
					"stock_uom",
					"warehouse",
				],
			),
			(
				self.query("movements", "records"),
				"records",
				[
					"document",
					"kind",
					"date",
					"available_effect",
					"stock_uom",
					"item_count",
					"source",
					"destination",
					"docstatus",
				],
			),
			(
				self.query("movements", "lines"),
				"history",
				[
					"item_code",
					"item_name",
					"document",
					"kind",
					"date",
					"docstatus",
					"source_warehouse",
					"destination_warehouse",
					"available_effect",
					"stock_uom",
				],
			),
			(
				self.query("loans", "records"),
				"records",
				[
					"document",
					"borrower",
					"activity_title",
					"date",
					"outstanding",
					"borrowed",
					"stock_uom",
					"item_count",
					"source",
					"destination",
					"loan_status",
				],
			),
			(
				self.query("loans", "items"),
				"loan_history",
				[
					"item_code",
					"item_name",
					"document",
					"kind",
					"date",
					"activity_title",
					"borrower",
					"docstatus",
					"source_warehouse",
					"destination_warehouse",
					"outstanding",
					"quantity",
					"stock_uom",
				],
			),
			(
				self.query("counts", "records"),
				"records",
				[
					"document",
					"date",
					"time",
					"docstatus",
					"count_scope",
					"item_count",
					"gains",
					"losses",
					"net",
					"baseline_qty",
					"actual_qty",
					"difference",
					"stock_uom",
				],
			),
			(
				self.query("counts", "lines"),
				"lines",
				[
					"item_name",
					"item_code",
					"warehouse",
					"batch_no",
					"baseline_qty",
					"actual_qty",
					"difference",
					"stock_uom",
					"reason",
				],
			),
			(
				self.query("activities", "metadata"),
				"metadata",
				[
					"title",
					"activity_type_label",
					"status",
					"start_date",
					"end_date",
					"document_count",
					"description",
					"notes",
				],
			),
		]
		with patch("temple_inventory.exports._enqueue"):
			for query, table, columns in cases:
				with self.subTest(dataset=query["dataset"], table=table):
					result = exports.request_export(query, table, columns=columns)
					request = frappe.get_doc("Inventory Export Request", result["name"])
					self.assertEqual(json.loads(request.columns_json), columns)

	def test_browse_export_headers_match_page_column_labels(self):
		contracts = {
			("inventory", "summary"): {"stock_qty": "合计", "stock_uom": "库存单位"},
			("expiry", "batches"): {
				"stock_qty": "库存数量",
				"warehouse": "位置",
				"warehouse_path": "完整位置",
				"warehouse_native_name": "ERPNext 库位",
				"expiry_date": "到期日期",
				"remaining_days": "剩余天数",
				"expiry_state": "到期状态",
			},
			("inventory", "locations"): {
				"warehouse_label": "库位",
				"batch_no": "批次编号",
				"expiry_date": "到期日期",
				"remaining_days": "剩余天数",
				"expiry_state": "到期状态",
				"warehouse_path": "完整位置",
				"warehouse_native_name": "ERPNext 库位",
			},
			("movements", "records"): {
				"document": "编号",
				"source": "来源位置",
				"destination": "去向位置",
				"source_warehouse_path": "来源完整位置",
				"destination_warehouse_native_name": "去向 ERPNext 库位",
				"time": "时间",
				"available_effect": "可用变化",
				"original_loan": "原借出",
			},
			("loans", "records"): {
				"document": "借出单号",
				"date": "借出日期",
				"borrowed": "借出数量",
				"outstanding": "未归还数量",
				"loan_status": "借用状态",
				"time": "时间",
				"docstatus": "状态",
				"returned": "已归还",
				"damaged": "损坏",
				"lost": "遗失",
			},
			("loans", "loan_history"): {
				"document": "编号",
				"date": "日期",
				"kind": "类型",
				"source_warehouse": "来源位置",
				"destination_warehouse": "去向位置",
				"quantity": "数量",
				"outstanding": "未归还数量",
				"stock_uom": "单位",
				"time": "时间",
				"original_loan": "原借出",
			},
			("counts", "lines"): {"difference": "差异量", "warehouse_path": "完整位置"},
		}
		for (dataset, table), fields in contracts.items():
			for field, expected in fields.items():
				self.assertEqual(export_label(dataset, field, table), expected)

	def test_stock_browse_returns_total_for_rows_and_location_summaries(self):
		self.receipt(12)
		result = browse._inventory_page(self.query(), 0, 30)
		row = next(value for value in result["rows"] if value["code"] == self.item)
		self.assertEqual(row["stock_qty"], 12)
		self.assertEqual(sum(location["stock_qty"] for location in row["locations"]), 12)

	def test_grouped_xlsx_writes_labeled_headers_subtotals_and_outline_levels(self):
		self.receipt(12)
		query = self.query("movements", "lines", groupBy=["item"])
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(
				query,
				"lines",
				columns=["item_name", "stock_qty", "stock_uom", "available_effect"],
				grouped=True,
			)
		request = frappe.get_doc("Inventory Export Request", result["name"])
		self.assertTrue(request.grouped)
		with Projection(query, refresh=True) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			path = Path(directory) / "grouped.xlsx"
			self.assertEqual(exports._write(projection, request, path), 1)
			workbook = load_workbook(path, data_only=True)
			worksheet = workbook["库存明细"]
			self.assertEqual(worksheet.cell(2, 1).value, "Native test item (1)")
			self.assertEqual(worksheet.cell(2, 2).value, 12)
			self.assertEqual(worksheet.cell(2, 3).value, uom_label("Nos"))
			self.assertEqual(worksheet.cell(2, 4).value, 12)
			self.assertEqual(worksheet.row_dimensions[2].outline_level, 0)
			self.assertEqual(worksheet.row_dimensions[3].outline_level, 1)
			self.assertFalse(worksheet.row_dimensions[3].hidden)
			self.assertEqual(worksheet.cell(4, 1).value, f"合计可用变化（{uom_label('Nos')}） · 1 行")
			self.assertEqual(worksheet.cell(4, 2).value, 12)
			self.assertEqual(worksheet.cell(4, 4).value, 12)
			self.assertEqual(worksheet.cell(2, 1).fill.fgColor.rgb[-6:], "D9D9D9")
			workbook.close()
		if not frappe.db.exists("UOM", "Box"):
			frappe.get_doc({"doctype": "UOM", "uom_name": "Box"}).insert()
		other = frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": "TI-BOX-" + self.token,
				"item_name": "Native test item in boxes",
				"stock_uom": "Box",
				"item_group": "All Item Groups",
				"is_stock_item": 1,
			}
		).insert()
		data = self.movement(quantity=3)
		data["items"][0].update({"item_code": other.name, "uom": "Box", "conversion_factor": 1})
		self.submit(self.create(data))
		category_query = self.query("movements", "lines", groupBy=["category"])
		category_query["scope"]["item"] = ""
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(
				category_query,
				"lines",
				columns=["item_name", "stock_qty", "stock_uom", "available_effect"],
				grouped=True,
			)
		request = frappe.get_doc("Inventory Export Request", result["name"])
		with (
			Projection(category_query, refresh=True) as projection,
			tempfile.TemporaryDirectory() as directory,
		):
			projection.build()
			path = Path(directory) / "category.xlsx"
			self.assertEqual(exports._write(projection, request, path), 2)
			workbook = load_workbook(path, data_only=True)
			worksheet = workbook.worksheets[0]
			self.assertEqual(worksheet.cell(2, 1).value, "无分类 (2)")
			self.assertEqual(worksheet.cell(3, 1).value, f"小计 · {uom_label('Box')}")
			self.assertEqual(worksheet.cell(3, 2).value, 3)
			self.assertEqual(worksheet.cell(3, 3).value, uom_label("Box"))
			self.assertEqual(worksheet.cell(3, 4).value, 3)
			self.assertEqual(worksheet.cell(4, 1).value, f"小计 · {uom_label('Nos')}")
			self.assertEqual(worksheet.cell(4, 2).value, 12)
			self.assertEqual(worksheet.cell(4, 3).value, uom_label("Nos"))
			self.assertEqual(worksheet.cell(4, 4).value, 12)
			self.assertEqual(worksheet.row_dimensions[2].outline_level, 0)
			workbook.close()

	def test_grouped_csv_rejected_empty_grouping_clears_flag_and_file_columns_drop(self):
		query = self.query("movements", "lines")
		with self.assertRaises(frappe.ValidationError):
			exports.request_export(query, "lines", export_format="csv", grouped=True)
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(
				query,
				"lines",
				columns=["image_file", "attachments"],
				grouped=True,
			)
		request = frappe.get_doc("Inventory Export Request", result["name"])
		self.assertFalse(request.grouped)
		self.assertNotIn("image_file", json.loads(request.columns_json))
		self.assertNotIn("attachments", json.loads(request.columns_json))

	def test_zip_export_includes_primary_image_attachments_hyperlink_and_manifest_files(self):
		# Files stay on disk after rollback, and Frappe renames a clashing upload.
		notes = f"notes-{uuid4().hex[:8]}"
		self.receipt(12)
		primary = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": "primary.webp",
				"content": b"image bytes",
				"is_private": 1,
				"attached_to_doctype": "Item",
				"attached_to_name": self.item,
				"attached_to_field": "image",
			}
		).insert()
		item = frappe.get_doc("Item", self.item)
		item.image = primary.file_url
		item.save()
		for _ in range(2):
			frappe.get_doc(
				{
					"doctype": "File",
					"file_name": f"{notes}.txt",
					"content": b"attachment",
					"is_private": 1,
					"attached_to_doctype": "Item",
					"attached_to_name": self.item,
				}
			).insert()
		record = self.create(self.movement(quantity=3))
		query = self.query("movements", "lines")
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(
				query,
				"summary",
				columns=["item_code", "image_file", "attachments"],
				include_files="all",
			)
		request = frappe.get_doc("Inventory Export Request", result["name"])
		self.assertTrue(request.file_label.endswith(".zip"))
		with Projection(query, refresh=True) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			path = Path(directory) / request.file_label
			exports._write(projection, request, path)
			self.assertEqual(
				projection.export_file_ids,
				{file.name for file in native_files.attachments("Item", self.item)},
			)
			with zipfile.ZipFile(path) as archive:
				names = set(archive.namelist())
				self.assertIn(f"images/{self.item}.webp", names)
				self.assertIn(f"attachments/items/{self.item}/{notes}.txt", names)
				self.assertIn(f"attachments/items/{self.item}/{notes} (2).txt", names)
				book = load_workbook(io.BytesIO(archive.read(request.file_label.rsplit(".", 1)[0] + ".xlsx")))
				worksheet = book["汇总"]
				self.assertEqual(worksheet.cell(2, 2).value, f"images/{self.item}.webp")
				self.assertIn(
					self.item,
					worksheet.cell(2, 2).hyperlink.target if worksheet.cell(2, 2).hyperlink else "",
				)
				self.assertEqual(worksheet.cell(2, 3).value, f"attachments/items/{self.item}/")
				book.close()
			frappe.get_doc(
				{
					"doctype": "File",
					"file_name": f"{notes}-record.txt",
					"content": b"record attachment",
					"is_private": 1,
					"attached_to_doctype": "Stock Entry",
					"attached_to_name": record.name,
				}
			).insert()
			record_query = self.query("movements", "lines", record=record.name, statuses=[0])
			with patch("temple_inventory.exports._enqueue"):
				record_request = frappe.get_doc(
					"Inventory Export Request",
					exports.request_export(
						record_query,
						"lines",
						columns=["document", "attachments"],
						include_files="all",
					)["name"],
				)
			record_path = Path(directory) / record_request.file_label
			with Projection(record_query, refresh=True) as record_projection:
				record_projection.build()
				exports._write(record_projection, record_request, record_path)
				with zipfile.ZipFile(record_path) as archive:
					self.assertIn(f"attachments/records/{record.name}/{notes}-record.txt", archive.namelist())
				with patch(
					"temple_inventory.item_images._authorized_original", side_effect=frappe.PermissionError
				):
					with self.assertRaises(frappe.PermissionError):
						exports._write(record_projection, record_request, record_path)

	def test_history_sort_matches_browse_and_retry_preserves_export_options(self):
		self.receipt(12)
		query = self.query("movements", "lines", groupBy=["item"])
		expected = browse.movement_history(query, sort_by="document", sort="asc")["rows"]
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(
				query,
				"history",
				columns=["document", "attachments"],
				grouped=True,
				include_files="images",
				sort_by="document",
				sort="asc",
			)
		request = frappe.get_doc("Inventory Export Request", result["name"])
		self.assertTrue(request.grouped)
		self.assertEqual(request.include_files, "images")
		self.assertEqual(request.export_sort_by, "document")
		self.assertEqual(request.export_sort, "asc")
		with Projection(query, refresh=True) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			path = Path(directory) / request.file_label
			exports._write(projection, request, path)
			with zipfile.ZipFile(path) as archive:
				workbook = load_workbook(
					io.BytesIO(archive.read(request.file_label.rsplit(".", 1)[0] + ".xlsx")), data_only=True
				)
				worksheet = workbook["流动明细"]
				expected_names = {row["document"] for row in expected}
				data_rows = [row for row in list(worksheet.values)[1:] if row[0] in expected_names]
				values = [row[0] for row in data_rows]
				self.assertEqual(values, [row["document"] for row in expected])
				self.assertTrue(all(not row[1] for row in data_rows))
				workbook.close()
		request.status = "Failed"
		exports._save(request)
		with patch("temple_inventory.exports._enqueue"):
			retried = exports.retry_export(request.name, str(request.modified))
		self.assertEqual(retried["attempt"], 2)
		self.assertTrue(retried["grouped"])
		self.assertEqual(retried["include_files"], "images")
		self.assertEqual(request.export_sort_by, "document")
		self.assertEqual(request.export_sort, "asc")

	def test_export_totals_are_per_unit_and_expiry_states_are_date_accurate(self):
		stream = io.StringIO()
		writer = csv.writer(stream)
		count = exports._write_rows(
			writer,
			[
				{"item_name": "甲", "stock_qty": 2, "stock_uom": "Nos"},
				{"item_name": "乙", "stock_qty": 3, "stock_uom": "Nos"},
				{"item_name": "丙", "stock_qty": 4, "stock_uom": "Box"},
			],
			["item_name", "stock_qty"],
			"batches",
			frappe._dict(query={"dataset": "inventory"}),
		)
		self.assertEqual(count, 3)
		values = list(csv.reader(io.StringIO(stream.getvalue())))
		self.assertEqual([row[1] for row in values[-2:]], ["5", "4"])
		self.assertIn("合计", values[-2][0])
		self.assertEqual(exports._display_cell({"expiry_date": "2000-01-01"}, "expiry_state"), "已过期")
		self.assertEqual(exports._display_cell({"expiry_date": "2999-01-01"}, "expiry_state"), "未过期")
		self.assertEqual(exports._display_cell({}, "expiry_state"), "无效期")

	def test_grouped_movement_export_uses_flat_history_rows_and_readable_cells(self):
		self.receipt(12)
		query = self.query("movements", "summary", groupBy=["item"])
		with Projection(query, refresh=True) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			self.assertEqual(projection.count("history"), browse.movement_history(query)["matched_count"])
			path = Path(directory) / "history.csv"
			count = exports._write(
				projection,
				frappe._dict(
					tables_json='["history"]',
					columns_json=json.dumps(
						[
							"kind",
							"docstatus",
							"destination_warehouse",
							"stock_qty",
							"stock_uom",
							"available_effect",
						]
					),
					export_format="csv",
				),
				path,
			)
			with path.open(encoding="utf-8-sig", newline="") as stream:
				rows = list(csv.reader(stream))
			self.assertEqual(count, 1)
			self.assertEqual(rows[0], ["类型", "状态", "去向位置", "数量", "单位", "可用变化"])
			self.assertEqual(rows[1][0], "入库")
			self.assertEqual(rows[1][1], "已提交")
			self.assertEqual(rows[1][2], projection.context.label(self.a))
			self.assertNotEqual(rows[1][2], self.a)
			self.assertEqual(rows[1][4], "个")
			# Quantity stays separate; per-unit totals sum the available-effect column.
			self.assertEqual(rows[1][3], "12")
			self.assertEqual(rows[-1][0], f"合计可用变化（{uom_label('Nos')}） · 1 行")
			self.assertEqual(rows[-1][3], "")
			self.assertEqual(rows[-1][5], "12")
			self.assertTrue(all(not any(isinstance(cell, (list, dict)) for cell in row) for row in rows))

	def test_loan_history_export_lists_only_outstanding_submitted_loans(self):
		partly = self.loan(5, borrower="部分归还")
		self.submit(self.settlement(partly, 2))
		settled = self.loan(3, borrower="全部归还")
		self.submit(self.settlement(settled, 3))
		query = self.query("loans", "items")
		with Projection(query, refresh=True) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			self.assertEqual(len(list(projection.rows("loan_history"))), 1)
			path = Path(directory) / "loans.csv"
			exports._write(
				projection,
				frappe._dict(
					tables_json='["loan_history"]',
					columns_json=json.dumps(
						[
							"document",
							"kind",
							"item_name",
							"borrower",
							"activity_title",
							"source_warehouse",
							"quantity",
							"stock_uom",
						]
					),
					export_format="csv",
				),
				path,
			)
			with path.open(encoding="utf-8-sig", newline="") as stream:
				rows = list(csv.reader(stream))
		self.assertEqual(
			rows[0],
			["编号", "类型", "物品名称", "借用方", "活动", "来源位置", "数量", "单位"],
		)
		self.assertEqual([row[0] for row in rows[1:-1]], [partly.name])
		self.assertEqual(rows[1][3:6], ["部分归还", self.activity.title, projection.context.label(self.a)])
		self.assertEqual(float(rows[1][6]), 5)
		self.assertTrue(rows[-1][0].startswith("合计"))
		inclusive_query = self.query("loans", "items", includeZero=True)
		with (
			Projection(inclusive_query, refresh=True) as inclusive_projection,
			tempfile.TemporaryDirectory() as directory,
		):
			inclusive_projection.build()
			inclusive_path = Path(directory) / "loans-including-zero.csv"
			exports._write(
				inclusive_projection,
				frappe._dict(
					tables_json='["loan_history"]',
					columns_json=json.dumps(["document", "quantity", "stock_uom"]),
					export_format="csv",
				),
				inclusive_path,
			)
			with inclusive_path.open(encoding="utf-8-sig", newline="") as stream:
				inclusive_rows = list(csv.reader(stream))
		self.assertEqual(
			{row[0] for row in inclusive_rows[1:-1]},
			{partly.name, settled.name},
		)

	def test_all_scope_loan_export_contains_loan_and_settlement_lines(self):
		partly = self.loan(5, borrower="部分归还")
		partly_return = self.submit(self.settlement(partly, 2))
		settled = self.loan(3, borrower="全部归还")
		settled_return = self.submit(self.settlement(settled, 3))
		query = self.query("loans", "items", loanScope="all")
		with Projection(query, refresh=True) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			path = Path(directory) / "loans-all.csv"
			exports._write(
				projection,
				frappe._dict(
					tables_json='["loan_history"]',
					columns_json=json.dumps(["document", "kind", "item_name", "quantity", "stock_uom"]),
					export_format="csv",
				),
				path,
			)
			with path.open(encoding="utf-8-sig", newline="") as stream:
				rows = list(csv.reader(stream))
		self.assertEqual(rows[0], ["编号", "类型", "物品名称", "数量", "单位"])
		self.assertEqual({row[1] for row in rows[1:-1]}, {"借出", "归还"})
		self.assertEqual(
			{row[0] for row in rows[1:-1]},
			{partly.name, partly_return.name, settled.name, settled_return.name},
		)
		self.assertEqual(len(rows[1:-1]), 4)

	def test_export_api_rejects_internal_tables_and_disallowed_columns_and_saves_summary(self):
		self.receipt(12)
		query = self.query("movements", "summary")
		for table in ("grouped", "items", "files"):
			with self.assertRaises(frappe.ValidationError):
				exports.request_export(query, table)
		with self.assertRaises(frappe.ValidationError):
			exports.request_export(query, "summary", columns=["document_id"])
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(
				query,
				"summary",
				columns=["item_name", "stock_qty"],
				summary={"title": "月度库存", "filters": ["位置：供灯区"]},
			)
		request = frappe.get_doc("Inventory Export Request", result["name"])
		self.assertEqual(result["table"], "summary")
		self.assertEqual(json.loads(request.columns_json), ["item_name", "stock_qty"])
		self.assertEqual(json.loads(request.summary_json), {"title": "月度库存", "filters": ["位置：供灯区"]})
		self.assertTrue(request.file_label.startswith("月度库存-全部时间-"))
		self.assertTrue(request.file_label.endswith(".xlsx"))
		with Projection(query, refresh=True) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			path = Path(directory) / request.file_label
			exports._write(projection, request, path)
			workbook = load_workbook(path, read_only=True, data_only=True)
			self.assertEqual(workbook.sheetnames[0], "汇总")
			self.assertEqual(workbook.sheetnames[1], "导出说明")
			info_rows = list(workbook["导出说明"].values)
			workbook.close()
			self.assertIn(("页面", "月度库存"), info_rows)
			self.assertIn(("筛选", "位置：供灯区"), info_rows)
			self.assertIn(("分组", "无"), info_rows)

	def test_count_record_export_names_checked_locations_without_outside_destination(self):
		self.receipt(10)
		draft = counts.save_count(str(uuid4()), self.count_data(8, "Reviewed"))
		completed = native_api.submit_document("Stock Reconciliation", draft.name, draft.version)
		query = self.query("movements", "records", record=completed.name)
		with Projection(query, refresh=True) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			path = Path(directory) / "count.csv"
			exports._write(
				projection,
				frappe._dict(
					tables_json='["records"]',
					columns_json=json.dumps(["kind", "source", "destination"]),
					export_format="csv",
				),
				path,
			)
			with path.open(encoding="utf-8-sig", newline="") as stream:
				rows = list(csv.reader(stream))
		self.assertEqual(rows[1][0], "盘点调整")
		self.assertTrue(rows[1][1])
		self.assertEqual(rows[1][2], "")

	def test_count_record_export_keeps_quantity_cells_numeric_with_unit(self):
		self.receipt(10)
		draft = counts.save_count(str(uuid4()), self.count_data(8, "Reviewed"))
		completed = native_api.submit_document("Stock Reconciliation", draft.name, draft.version)
		query = self.query("counts", "records", record=completed.name, statuses=[1])
		columns = ["document", "gains", "losses", "net", "stock_uom"]
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(query, "records", columns=columns, export_format="csv")
		request = frappe.get_doc("Inventory Export Request", result["name"])
		with Projection(query, refresh=True) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			path = Path(directory) / "count-totals.csv"
			exports._write(projection, request, path)
			with path.open(encoding="utf-8-sig", newline="") as stream:
				rows = list(csv.reader(stream))
		self.assertEqual(rows[0], ["盘点编号", "增加", "减少", "净差异", "单位"])
		self.assertEqual(rows[1], [completed.name, "0", "-2", "-2", "个"])
		self.assertEqual(rows[2][1:], ["0", "-2", "-2", "个"])

	def test_records_export_collapses_units_to_one_native_document_row_with_locations(self):
		doc = self.submit(self.create(self.movement(quantity=2)))
		query = self.query("movements", "records", record=doc.name, statuses=[1])
		with Projection(query, refresh=True) as projection:
			projection.build()
			rows = list(exports._export_rows(projection, "records"))
		self.assertEqual(len(rows), 1)
		self.assertEqual(rows[0]["document"], doc.name)
		self.assertTrue(rows[0]["source"])
		self.assertTrue(rows[0]["destination"])
		self.assertNotIn(self.a, rows[0]["source"] + rows[0]["destination"])

	def test_archive_names_number_repeated_file_names(self):
		used = set()
		names = [exports._unique_archive_name(name, used) for name in ("notes.txt", "NOTES.txt", "notes.txt")]
		self.assertEqual(names, ["notes.txt", "NOTES (2).txt", "notes (3).txt"])

	def test_record_export_adds_item_counts_across_unit_rows(self):
		unit_rows = [
			{
				"doctype": "Draft",
				"document": "STE-1",
				"stock_uom": "Kg",
				"item_count": 1,
				"available_effect": 2,
			},
			{
				"doctype": "Draft",
				"document": "STE-1",
				"stock_uom": "Nos",
				"item_count": 65,
				"available_effect": 9,
			},
		]
		projection = frappe._dict(query={"dataset": "movements"}, rows=lambda _table: iter(unit_rows))
		rows = list(exports._export_rows(projection, "records"))
		self.assertEqual(len(rows), 2)
		self.assertEqual(sum(row["item_count"] for row in rows), 66)

	def test_count_record_export_sums_gains_and_losses_by_line_not_net(self):
		unit_rows = [
			{"doctype": "Stock Reconciliation", "document": "SR-1", "stock_uom": "Nos", "difference": 2},
		]
		projection = frappe._dict(
			query={"dataset": "counts"}, context=None, rows=lambda _table: iter(unit_rows)
		)
		operation = {"gains": [{"uom": "Nos", "qty": 5}], "losses": [{"uom": "Nos", "qty": -3}]}
		with (
			patch("temple_inventory.report_sources.whole_operation", return_value=operation),
			patch.object(exports, "_export_record_row", side_effect=lambda *args: args[3]),
		):
			rows = list(exports._export_rows(projection, "records"))
		self.assertEqual(
			rows[0]["_units"]["Nos"], {"difference": 2.0, "gains": 5.0, "losses": -3.0, "net": 2.0}
		)

	def test_line_export_leaves_a_count_adjustments_unnamed_end_blank(self):
		count_line = {"kind": "Reconcile", "source_warehouse": self.a, "destination_warehouse": ""}
		receipt_line = {"kind": "Receive", "source_warehouse": "", "destination_warehouse": self.a}
		self.assertEqual(exports._display_cell(count_line, "destination_warehouse"), "")
		self.assertEqual(exports._display_cell(receipt_line, "source_warehouse"), "外部")

	def test_mixed_unit_record_exports_numeric_quantities_with_units(self):
		class Context:
			@staticmethod
			def label(name):
				return {"source-warehouse": "来源库位", "destination-warehouse": "去向库位"}[name]

			@staticmethod
			def location_fields(name):
				return {
					"source-warehouse": {
						"warehouse_path": "来源路径",
						"warehouse_native_name": "ERP 来源",
					},
					"destination-warehouse": {
						"warehouse_path": "去向路径",
						"warehouse_native_name": "ERP 去向",
					},
				}.get(name, {})

		projection = frappe._dict(query={"dataset": "movements"}, context=Context())
		row = exports._export_record_row(
			projection,
			lambda *_args: {
				"branches": [{"source": "source-warehouse", "destination": "destination-warehouse"}]
			},
			{},
			{
				"doctype": "Stock Entry",
				"document": "STE-1",
				"_units": {
					"Nos": {"stock_qty": 12, "available_effect": -12},
					"Box": {"stock_qty": 3, "available_effect": 3},
				},
			},
		)
		self.assertIsNone(row["stock_qty"])
		self.assertEqual(row["stock_uom"], "")
		self.assertIsNone(row["available_effect"])
		self.assertEqual(row["source"], "来源库位")
		self.assertEqual(row["destination"], "去向库位")

	def test_mixed_unit_record_rows_split_into_numeric_unit_rows(self):
		unit_rows = [
			{
				"doctype": "Draft",
				"document": "STE-1",
				"stock_uom": "Nos",
				"item_count": 1,
				"available_effect": -12,
				"stock_qty": 12,
			},
			{
				"doctype": "Draft",
				"document": "STE-1",
				"stock_uom": "Box",
				"item_count": 1,
				"available_effect": 3,
				"stock_qty": 3,
			},
		]
		projection = frappe._dict(
			query={"dataset": "movements"},
			rows=lambda _table: iter(unit_rows),
		)
		rows = list(exports._export_rows(projection, "records"))
		self.assertEqual(len(rows), 2)
		self.assertEqual(
			{(row["stock_uom"], row["stock_qty"], row["available_effect"]) for row in rows},
			{("Nos", 12.0, -12.0), ("Box", 3.0, 3.0)},
		)

	def test_movement_history_export_includes_original_loan_references(self):
		row = {"doctype": "Stock Entry", "document": "STE-1", "item_code": self.item}
		projection = frappe._dict(
			query={"dataset": "movements"},
			context=object(),
			rows=lambda _table: iter([row]),
		)
		with patch(
			"temple_inventory.report_sources.whole_operation",
			return_value={"original_loans": ["LOAN-1", "LOAN-2"]},
		):
			exported = list(exports._export_rows(projection, "history"))
		self.assertEqual(exported[0]["original_loan"], "LOAN-1；LOAN-2")

	def test_refinement_display_labels_preserve_native_identity_and_export_audit_values(self):
		room = frappe.get_doc("Warehouse", self.room_a)
		room.ti_label = "供灯-用品"
		room.save()
		doc = self.submit(self.create(self.movement(quantity=2)))
		query = self.query("movements", "lines", record=doc.name, statuses=[1])
		with Projection(query) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			row = next(projection.rows("lines"))
			self.assertEqual(row["docstatus"], 1)
			self.assertEqual(row["docstatus_label"], "已提交")
			self.assertEqual(row["destination_warehouse"], self.a)
			self.assertEqual(row["destination_warehouse_label"], "供灯-用品")
			self.assertEqual(
				row["destination_warehouse_path"],
				projection.context.location_fields(self.a)["warehouse_path"],
			)
			self.assertEqual(
				row["destination_warehouse_native_name"],
				frappe.db.get_value("Warehouse", self.a, "warehouse_name"),
			)
			self.assertEqual(projection.context.label("inaccessible-warehouse"), "")
			path = Path(directory) / "labels.csv"
			exports._write(projection, frappe._dict(tables_json='["lines"]', export_format="csv"), path)
			with path.open(encoding="utf-8-sig", newline="") as stream:
				values = next(csv.DictReader(stream))
			self.assertEqual(values["状态"], "已提交")
		self.assertEqual(values["去向位置"], "供灯-用品")
		self.assertNotIn(self.a, values.values())

	def test_saved_opening_export_uses_native_movement_quantities_and_counts_opening_stock(self):
		doc = frappe.get_doc(
			{
				"doctype": "Stock Reconciliation",
				"company": self.company,
				"purpose": "Opening Stock",
				"expense_account": frappe.db.get_value(
					"Account", {"company": self.company, "account_name": "Temporary Opening"}, "name"
				),
				"posting_date": nowdate(),
				"items": [{"item_code": self.item, "warehouse": self.a, "qty": 7, "valuation_rate": 10}],
			}
		).insert()
		for status in (0, 1, 2):
			if status == 1:
				doc.submit()
			elif status == 2:
				doc.cancel()
			loaded = native_api.load_document(doc.doctype, doc.name)
			self.assertEqual(loaded.purpose, "Opening Stock")
			query = self.query(
				"movements", "records", record=doc.name, recordDoctype=doc.doctype, statuses=[status]
			)
			page = report_page(query, page_length=5)
			self.assertEqual(page["matched_count"], 1)
			self.assertEqual(page["rows"][0]["document"], doc.name)
			self.assertEqual(page["rows"][0]["kind"], "Opening")
			with Projection(query) as projection:
				projection.build()
				lines = list(projection.rows("lines"))
				self.assertEqual(len(lines), 1)
				self.assertEqual(lines[0]["stock_qty"], 7)
				self.assertEqual(lines[0]["actual_qty"], 7)
				self.assertEqual(lines[0]["doctype"], doc.doctype)
				self.assertEqual(lines[0]["line_name"], doc.items[0].name)
			count_page = report_page({**query, "dataset": "counts"}, page_length=5)
			self.assertEqual(count_page["matched_count"], 1)
			self.assertEqual(count_page["rows"][0]["document"], doc.name)
			self.assertEqual(count_page["rows"][0]["baseline_qty"], 0)
			self.assertEqual(count_page["rows"][0]["actual_qty"], 7)
			self.assertEqual(report_page({**query, "recordDoctype": "Stock Entry"})["matched_count"], 0)

	def test_xlsx_and_csv_keep_identical_selected_numeric_table_scope(self):
		self.receipt(12)
		query = self.query("movements", "lines")
		with tempfile.TemporaryDirectory() as directory, Projection(query) as projection:
			projection.build(related=True)
			request = frappe._dict(
				tables_json='["lines"]',
				columns_json='["stock_qty"]',
				owner="Administrator",
				export_format="csv",
			)
			csv_path = Path(directory) / "report.csv"
			csv_count = exports._write(projection, request, csv_path)
			with csv_path.open(encoding="utf-8-sig") as source:
				csv_rows = list(csv.DictReader(source))
			request.export_format = "xlsx"
			xlsx_path = Path(directory) / "report.xlsx"
			xlsx_count = exports._write(projection, request, xlsx_path)
			workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
			rows = list(workbook["库存明细"].values)
			workbook.close()
			self.assertEqual(csv_count, xlsx_count)
			self.assertEqual(len(csv_rows), len(rows) - 1)
			self.assertEqual(float(csv_rows[0]["库存数量"]), 12)
			self.assertEqual(rows[1][rows[0].index("库存数量")], 12)

	def test_large_native_draft_export_preserves_all_1005_allocations_and_stable_order(self):
		data = self.movement(quantity=1)
		data["items"] = [copy.deepcopy(data["items"][0]) for _ in range(1005)]
		doc = self.create(data)
		query = self.query(
			"movements", "lines", record=doc.name, statuses=[0], sortBy="item_name", sort="asc"
		)
		preview = report_page(query, page_length=1)
		self.assertEqual(preview["matched_count"], 1005)
		self.assertEqual(len(preview["rows"]), 1)
		with tempfile.TemporaryDirectory() as directory, Projection(query) as projection:
			projection.build()
			request = frappe._dict(
				tables_json='["lines"]',
				columns_json='["line_name", "stock_qty"]',
				owner="Administrator",
				export_format="csv",
			)
			csv_path = Path(directory) / "large.csv"
			self.assertEqual(exports._write(projection, request, csv_path), 1005)
			with csv_path.open(encoding="utf-8-sig") as source:
				csv_rows = list(csv.DictReader(source))
			request.export_format = "xlsx"
			xlsx_path = Path(directory) / "large.xlsx"
			self.assertEqual(exports._write(projection, request, xlsx_path), 1005)
			workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
			rows = iter(workbook["库存明细"].values)
			columns = next(rows)
			ids = [
				row[columns.index("明细编号")]
				for row in rows
				if row[columns.index("明细编号")]
				and not str(row[columns.index("明细编号")]).startswith("合计")
			]
			workbook.close()
			self.assertEqual(
				ids,
				[
					row["明细编号"]
					for row in csv_rows
					if row["明细编号"] and not row["明细编号"].startswith("合计")
				],
			)
			self.assertEqual(len(set(ids)), 1005)
			self.assertEqual(
				{row["明细编号"] for row in csv_rows if not row["明细编号"].startswith("合计")},
				{row.name for row in doc["items"]},
			)
			self.assertEqual(
				sum(float(row["库存数量"]) for row in csv_rows if not row["明细编号"].startswith("合计")),
				1005,
			)

	def test_export_request_is_owned_queued_and_failed_retry_retains_scope(self):
		self.receipt()
		query = self.query()
		with patch("temple_inventory.exports._enqueue") as enqueue:
			result = exports.request_export(query, "summary", export_format="xlsx")
			self.assertEqual(result["status"], "Queued")
			enqueue.assert_called_once()
			doc = frappe.get_doc("Inventory Export Request", result["name"])
			self.assertEqual(doc.owner, "Administrator")
			doc.status = "Failed"
			exports._save(doc)
			retried = exports.retry_export(doc.name, str(doc.modified))
			self.assertEqual(retried["attempt"], 2)
			self.assertEqual(json.loads(doc.query_json), query)
			self.assertEqual(json.loads(doc.tables_json), ["summary"])
		with self.assertRaises(frappe.ValidationError):
			exports.request_export(query, "summary,locations", export_format="csv")

	def test_export_private_download_rechecks_current_warehouse_and_request_owner(self):
		from frappe.cache_manager import clear_defaults_cache

		self.receipt()
		user = frappe.get_doc(
			{
				"doctype": "User",
				"email": f"ti-export-{self.token}@example.test",
				"first_name": "Export manager",
				"send_welcome_email": 0,
				"user_type": "System User",
				"roles": [{"role": "System Manager"}, {"role": "Stock Manager"}, {"role": "Stock User"}],
			}
		).insert()
		self.addCleanup(clear_defaults_cache, user.name)
		frappe.set_user(user.name)
		query = self.query()
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(query, "summary", export_format="csv")
		request = frappe.get_doc("Inventory Export Request", result["name"])
		request.status = "Running"
		exports._save(request)
		with tempfile.TemporaryDirectory() as directory, Projection(query) as projection:
			projection.build()
			output = Path(directory) / "export.csv"
			manifest = Path(directory) / "access.jsonl"
			exports._write(projection, request, output)
			exports._manifest(projection, manifest)
			request.output_file = exports._private_file(output, request, "export.csv")
			request.access_manifest = exports._private_file(manifest, request, "access.jsonl")
		for field in ("output_file", "access_manifest"):
			file = frappe.get_doc("File", request.get(field))
			self.addCleanup(Path(file.get_full_path()).unlink, missing_ok=True)
		request.status = "Completed"
		exports._save(request)
		with patch.object(exports, "validate_download", wraps=exports.validate_download) as validate:
			response = exports.download_export(request.name)
			validate.assert_called_once()
		self.assertEqual(response.status_code, 200)
		self.assertIn("attachment;", response.headers["Content-Disposition"])
		response.close()
		output_file = frappe.get_doc("File", request.output_file)
		self.assertTrue(output_file.is_downloadable())
		self.assertFalse(frappe.get_doc("File", request.access_manifest).is_downloadable())
		frappe.set_user("Administrator")
		with self.assertRaises(frappe.PermissionError):
			exports.export_status(request.name)
		frappe.get_doc(
			{"doctype": "User Permission", "user": user.name, "allow": "Warehouse", "for_value": self.b}
		).insert()
		frappe.clear_cache(user=user.name)
		frappe.set_user(user.name)
		with self.assertRaises(frappe.PermissionError):
			exports.validate_download(request)
		self.assertFalse(output_file.is_downloadable())

	def test_export_manifest_rechecks_names_in_bounded_doctype_batches(self):
		from types import SimpleNamespace

		names = [f"ITEM-{index}" for index in range(401)]
		with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8") as manifest:
			for name in names:
				manifest.write(json.dumps(["Item", name]) + "\n")
			manifest.flush()
			with (
				patch("temple_inventory.exports.require_app"),
				patch("temple_inventory.exports.frappe.db.get_value", return_value=1),
				patch(
					"temple_inventory.native_topology._settings",
					return_value=SimpleNamespace(company="Company"),
				),
				patch("temple_inventory.native_topology._visible_warehouses", return_value=set()),
				patch("temple_inventory.exports.frappe.get_doc") as get_doc,
				patch(
					"temple_inventory.exports.frappe.get_list",
					side_effect=lambda _doctype, filters, **_kwargs: filters["name"][1],
				) as get_list,
			):
				get_doc.return_value.check_permission = Mock()
				exports._check_manifest_path(manifest.name)
		self.assertEqual(get_list.call_count, 3)
		self.assertEqual(
			[len(call.kwargs["filters"]["name"][1]) for call in get_list.call_args_list], [200, 200, 1]
		)

	def test_movement_browse_always_includes_reconciliation_adjustments(self):
		self.receipt(10)
		draft = counts.save_count(str(uuid4()), self.count_data(8, "Reviewed"))
		completed = native_api.submit_document("Stock Reconciliation", draft.name, draft.version)
		query = self.query("movements", "lines", record=completed.name)
		self.assertEqual(report_page(query)["matched_count"], 1)
		self.assertEqual(report_page({**query, "includeAdjustments": False})["matched_count"], 1)
		self.assertEqual(report_page({**query, "includeAdjustments": "false"})["matched_count"], 1)
