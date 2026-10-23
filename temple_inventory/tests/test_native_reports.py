"""Report/export contracts exercised against real native stock and permissions."""

import copy
import csv
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

import frappe
from frappe.utils import nowdate
from openpyxl import load_workbook

from temple_inventory import activities, browse, catalog, counts, exports, native_api, preferences
from temple_inventory.projections import Projection, group_page, report_page
from temple_inventory.report_query import DEFINITIONS, validate_query
from temple_inventory.tests.test_native_documents import NativeDocumentTests


class NativeReportTests(NativeDocumentTests):
	def test_loan_item_contributions_keep_items_view_and_apply_independent_stable_sort(self):
		other = frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": "TI-LOAN-HISTORY-" + self.token,
				"item_name": "Second loan history item",
				"stock_uom": "Nos",
				"item_group": "All Item Groups",
				"is_stock_item": 1,
			}
		).insert()
		self.receipt(10)
		other_receipt = self.movement(quantity=10)
		other_receipt["items"][0]["item_code"] = other.name
		self.submit(self.create(other_receipt))
		data = self.movement("Loan", 3, self.a, self.loan_leaf, ti_borrower="History Sort")
		data["items"].append(dict(data["items"][0], item_code=other.name, qty=4))
		loan = self.submit(self.create(data))
		self.submit(self.settlement(loan, 3))
		query = self.query("loans", "items", loan=loan.name, statuses=[1], sortBy="item_name")
		page = browse.operation_contributions(
			query,
			item_code=self.item,
			history_scope="all",
			history_sort_by="document",
			history_sort="asc",
			page_length=1,
		)
		self.assertEqual(page["matched_count"], 2)
		self.assertEqual(page["next_start"], 1)
		self.assertEqual(page["rows"][0]["document"], loan.name)
		all_item_query = {**query, "scope": {**query["scope"], "item": ""}}
		all_items = browse.operation_contributions(
			all_item_query,
			history_scope="all",
			history_sort_by="document",
			history_sort="asc",
		)
		self.assertEqual(all_items["matched_count"], 3)
		self.assertEqual({row["item_code"] for row in all_items["rows"]}, {self.item, other.name})
		self.assertTrue(any(row["can_settle"] for row in all_items["rows"]))
		return_row = next(row for row in all_items["rows"] if row["kind"] == "Return")
		self.assertEqual(return_row["quantity"], -3)
		self.assertEqual(return_row["original_loans"], [loan.name])
		open_item_loan = next(
			row for row in all_items["rows"] if row["kind"] == "Loan" and row["item_code"] == other.name
		)
		self.assertEqual(open_item_loan["outstanding"], 4)
		self.assertEqual(open_item_loan["loan_status"], "Open")
		with self.assertRaises(frappe.ValidationError):
			browse.operation_contributions(
				query,
				item_code=self.item,
				history_scope="all",
				history_sort_by="unsupported",
			)
		outcome_query = {**query, "view": "outcomes"}
		with self.assertRaises(frappe.ValidationError):
			browse.operation_contributions(outcome_query, item_code=self.item)

	def test_grouped_native_record_leaves_include_whole_quantities_and_settlement_state(self):
		self.receipt(8)
		loan = self.submit(
			self.create(self.movement("Loan", 3, self.a, self.loan_leaf, ti_borrower="Grouped Loan"))
		)
		query = self.query("loans", "records", groupBy=["borrower"], loan=loan.name)
		root = group_page(query, parent_path=[])
		children = group_page(
			query,
			parent_path=["Grouped Loan"],
			parent_unit=root["rows"][0]["stock_uom"],
		)
		self.assertEqual(children["row_grain"], "record")
		self.assertEqual(len(children["rows"]), 1)
		row = children["rows"][0]
		self.assertEqual(row["row_grain"], "record")
		self.assertEqual(row["whole"]["quantities"], [{"uom": "Nos", "qty": 3.0}])
		self.assertEqual(row["whole"]["loan_status"], "Open")
		self.assertTrue(row["can_settle"])

	def test_location_presenter_labels_direct_leaves_and_disambiguates_duplicate_rooms(self):
		from temple_inventory.native_topology import present_locations

		settings = frappe._dict(root_warehouse="ROOT", physical_root_warehouse="PHYSICAL")
		rows = {
			"ROOT": frappe._dict(
				name="ROOT", warehouse_name="寺院仓库", ti_label="寺院仓库", parent_warehouse="", is_group=1
			),
			"PHYSICAL": frappe._dict(
				name="PHYSICAL",
				warehouse_name="实体库房",
				ti_label="实体库房",
				parent_warehouse="ROOT",
				is_group=1,
			),
			"AREA-1": frappe._dict(
				name="AREA-1",
				warehouse_name="T1",
				ti_label="第一寺院",
				ti_short_label="T1",
				parent_warehouse="PHYSICAL",
				warehouse_type="Area",
				is_group=1,
			),
			"AREA-2": frappe._dict(
				name="AREA-2",
				warehouse_name="T2",
				ti_label="第二寺院",
				ti_short_label="T2",
				parent_warehouse="PHYSICAL",
				warehouse_type="Area",
				is_group=1,
			),
			"ROOM-1": frappe._dict(
				name="ROOM-1",
				warehouse_name="T1 / A14",
				ti_label="A14",
				parent_warehouse="AREA-1",
				warehouse_type="Room",
				is_group=1,
			),
			"ROOM-2": frappe._dict(
				name="ROOM-2",
				warehouse_name="T2 / A14",
				ti_label="A14",
				parent_warehouse="AREA-2",
				warehouse_type="Room",
				is_group=1,
			),
			"DIRECT-1": frappe._dict(
				name="DIRECT-1",
				warehouse_name="T1 / A14 · 直接存放",
				ti_label="A14",
				parent_warehouse="ROOM-1",
				warehouse_type="",
				is_group=0,
				ti_direct_storage=1,
			),
			"DIRECT-2": frappe._dict(
				name="DIRECT-2",
				warehouse_name="T2 / A14 · 直接存放",
				ti_label="A14",
				parent_warehouse="ROOM-2",
				warehouse_type="",
				is_group=0,
				ti_direct_storage=1,
			),
		}
		presented = {row["name"]: row for row in present_locations(rows, settings=settings)}
		self.assertEqual(presented["DIRECT-1"]["label"], presented["ROOM-1"]["label"])
		self.assertIn("T1", presented["ROOM-1"]["label"])
		self.assertIn("T2", presented["ROOM-2"]["label"])
		self.assertTrue(presented["DIRECT-1"]["storable"])
		self.assertEqual(presented["ROOM-1"]["path"], "T1 / A14")

	def test_refinement_display_labels_preserve_native_identity_and_export_audit_values(self):
		room = frappe.get_doc("Warehouse", self.room_a)
		room.ti_label = "供灯-用品"
		room.save()
		doc = self.submit(self.create(self.movement(quantity=2)))
		query = self.query("movements", "lines", record=doc.name, statuses=[1])
		with Projection(query) as projection, tempfile.TemporaryDirectory() as directory:
			projection.build()
			row = next(projection.rows("lines"))
			self.assertEqual(row["docstatus"], 1)
			self.assertEqual(row["docstatus_label"], "已提交")
			self.assertEqual(row["destination_warehouse"], self.a)
			self.assertEqual(row["destination_warehouse_label"], "供灯-用品")
			self.assertEqual(projection.context.label("inaccessible-warehouse"), "")
			path = Path(directory) / "labels.csv"
			exports._write(projection, frappe._dict(tables_json='["lines"]', export_format="csv"), path)
			with path.open(encoding="utf-8-sig", newline="") as stream:
				values = next(csv.DictReader(stream))
			self.assertEqual(values["状态"], "已提交")
			self.assertEqual(values["状态值（ERPNext）"], "1")  # noqa: RUF001
			self.assertEqual(values["目标位置"], "供灯-用品")
			self.assertEqual(values["目标位置（ERPNext 名称）"], self.a)  # noqa: RUF001

	def test_repeated_allocations_authorize_item_images_once_per_report_context(self):
		from temple_inventory.native_files import attachments

		data = self.movement(quantity=1)
		data["items"] *= 12
		created = self.create(data)
		query = self.query("movements", "lines", record=created["name"], statuses=[0])
		with patch("temple_inventory.native_files.attachments", wraps=attachments) as image_reads:
			page = report_page(query)
		self.assertEqual(page["matched_count"], 12)
		self.assertEqual(image_reads.call_args_list, [unittest.mock.call("Item", self.item)])
		# The cache must not outlive the request/context or retain another account's
		# authorization. A second projection performs its own native File check.
		with patch("temple_inventory.native_files.attachments", wraps=attachments) as image_reads:
			self.assertEqual(report_page(query)["matched_count"], 12)
		self.assertEqual(image_reads.call_count, 1)

	def test_desk_multi_batch_count_retains_per_batch_baseline_zero_actual_and_cancelled_provenance(self):
		frappe.db.set_single_value("Stock Settings", "enable_serial_and_batch_no_for_item", 1)
		item = frappe.get_doc("Item", self.item)
		item.has_batch_no = 1
		item.save()
		batches = []
		for suffix in ("A", "B"):
			batch = catalog.create_batch(self.item, "TI-COUNT-MULTI-" + self.token + suffix)["name"]
			batches.append(batch)
			data = self.movement(quantity=5)
			data["items"][0].update({"batch_no": batch, "use_serial_batch_fields": 1})
			self.submit(self.create(data))
		posting = frappe.utils.now_datetime()
		bundle = frappe.get_doc(
			{
				"doctype": "Serial and Batch Bundle",
				"voucher_type": "Stock Reconciliation",
				"company": self.company,
				"item_code": self.item,
				"warehouse": self.a,
				"type_of_transaction": "Inward",
				"posting_datetime": posting,
				"entries": [{"batch_no": batches[0], "qty": 3}],
			}
		).insert()
		doc = frappe.get_doc(
			{
				"doctype": "Stock Reconciliation",
				"company": self.company,
				"purpose": "Stock Reconciliation",
				"posting_date": posting.date(),
				"posting_time": posting.time(),
				"set_posting_time": 1,
				"items": [
					{
						"item_code": self.item,
						"warehouse": self.a,
						"qty": 3,
						"serial_and_batch_bundle": bundle.name,
						"use_serial_batch_fields": 0,
						"reconcile_all_serial_batch": 1,
					}
				],
			}
		).insert()
		doc.submit()
		loaded = native_api.load_document("Stock Reconciliation", doc.name)
		self.assertEqual(
			{
				row["batch_no"]: (row["baseline_qty"], row["actual_qty"])
				for row in loaded["items"][0].recorded_count_allocations
			},
			{batches[0]: (5, 3), batches[1]: (5, 0)},
		)
		query = self.query("counts", "lines", record=doc.name, statuses=[1])
		before = report_page(query)
		self.assertEqual(before["matched_count"], 2)
		self.assertEqual(
			{
				row["batch_no"]: (row["baseline_qty"], row["actual_qty"], row["difference"])
				for row in before["rows"]
			},
			{batches[0]: (5, 3, -2), batches[1]: (5, 0, -5)},
		)
		movements = report_page(self.query("movements", "lines", record=doc.name, statuses=[1]))
		self.assertEqual(
			{row["batch_no"]: row["stock_qty"] for row in movements["rows"]}, {batches[0]: -2, batches[1]: -5}
		)
		cancelled = native_api.cancel_document("Stock Reconciliation", doc.name, str(doc.modified))
		self.assertEqual(
			cancelled["items"][0].recorded_count_allocations, loaded["items"][0].recorded_count_allocations
		)
		query["statuses"] = [2]
		after = report_page(query)
		fields = (
			"id",
			"line_name",
			"batch_no",
			"baseline_qty",
			"actual_qty",
			"difference",
			"bundle_line",
			"current_bundle_line",
		)
		self.assertEqual(
			[{key: row.get(key) for key in fields} for row in after["rows"]],
			[{key: row.get(key) for key in fields} for row in before["rows"]],
		)
		self.assertTrue(all(row["docstatus"] == 2 for row in after["rows"]))
		batch_class = type(frappe.get_doc("Batch", batches[1]))
		check = batch_class.check_permission

		def permission(batch_doc, *args, **kwargs):
			if batch_doc.name == batches[1]:
				raise frappe.PermissionError("Baseline Batch denied")
			return check(batch_doc, *args, **kwargs)

		with patch.object(batch_class, "check_permission", permission):
			with self.assertRaises(frappe.PermissionError):
				native_api.load_document("Stock Reconciliation", cancelled.name)
			self.assertEqual(report_page(query)["matched_count"], 0)

	def test_cancelled_batch_records_retain_native_allocation_ids_and_batch_permissions(self):
		frappe.db.set_single_value("Stock Settings", "enable_serial_and_batch_no_for_item", 1)
		item = frappe.get_doc("Item", self.item)
		item.has_batch_no = 1
		item.save()
		batch = catalog.create_batch(self.item, "TI-CANCEL-" + self.token)
		data = self.movement(quantity=4)
		data["items"][0].update({"batch_no": batch["name"], "use_serial_batch_fields": 1})
		completed = self.submit(self.create(data))
		query = self.query("movements", "lines", record=completed.name, statuses=[1])
		before = report_page(query)
		self.assertEqual(before["rows"][0]["batch_no"], batch["name"])
		cancelled = native_api.cancel_document("Stock Entry", completed.name, completed.version)
		query["statuses"] = [2]
		after = report_page(query)
		self.assertEqual(after["matched_count"], 1)
		self.assertEqual(after["rows"][0]["batch_no"], batch["name"])
		self.assertEqual(after["rows"][0]["bundle_line"], before["rows"][0]["bundle_line"])
		self.assertEqual(after["rows"][0]["stock_qty"], 4)
		self.assertEqual(cancelled["items"][0].recorded_bundle, completed["items"][0].recorded_bundle)
		self.assertEqual(
			cancelled["items"][0].recorded_batch_allocations, completed["items"][0].recorded_batch_allocations
		)
		batch_class = type(frappe.get_doc("Batch", batch["name"]))
		with patch.object(
			batch_class, "check_permission", side_effect=frappe.PermissionError("Batch denied")
		):
			with self.assertRaises(frappe.PermissionError):
				native_api.load_document("Stock Entry", cancelled.name)
			self.assertEqual(report_page(query)["matched_count"], 0)

	def query(self, dataset="inventory", view=None, **changes):
		value = {"dataset": dataset, "scope": {"item": self.item, "warehouses": [self.root]}}
		if view:
			value["view"] = view
		if view == "historical":
			value["scope"]["period"] = {"key": "this_month", "from": "", "to": ""}
		return validate_query({**value, **changes})

	def test_six_datasets_and_all_views_have_complete_native_page_contracts(self):
		self.receipt()
		for dataset, definition in DEFINITIONS.items():
			for view in definition["views"]:
				query = self.query(dataset, view)
				if view == "historical":
					query["scope"]["period"] = {"key": "this_month", "from": "", "to": ""}
				result = report_page(query, page_length=1)
				self.assertIn("matched_count", result)
				self.assertIn("snapshot", result)
				self.assertTrue(result["columns"])
				self.assertNotIn("id", result["columns"])
				for table in result["tables"]:
					self.assertNotIn("id", report_page(query, table=table["id"])["columns"])
				self.assertEqual(result["query"]["dataset"], dataset)
				field = {
					("inventory", "summary"): "item",
					("inventory", "locations"): "location",
					("inventory", "historical"): "item",
					("expiry", "batches"): "item",
					("movements", "summary"): "item",
					("movements", "records"): "kind",
					("movements", "lines"): "source",
					("loans", "records"): "borrower",
					("loans", "items"): "item",
					("loans", "outcomes"): "item",
					("loans", "settlements"): "item",
					("activities", "metadata"): "activityType",
					("activities", "records"): "kind",
					("activities", "loans"): "item",
					("counts", "records"): "kind",
					("counts", "lines"): "item",
				}[(dataset, view)]
				grouped = report_page({**query, "groupBy": [field]})
				self.assertEqual(grouped["table"], "grouped")
				self.assertNotIn("id", grouped["columns"])
				self.assertTrue(
					all("group_path" in row and row["group_depth"] == 1 for row in grouped["rows"])
				)
				if grouped["rows"] and field == "item":
					self.assertIn("group_label", grouped["columns"])
					self.assertTrue(grouped["rows"][0]["group_label"])

	def test_group_by_contract_rejects_retired_duplicate_unsupported_and_third_fields(self):
		query = self.query("inventory", "summary")
		self.assertEqual(
			validate_query({**query, "groupBy": ["category", "item"]})["groupBy"],
			["category", "item"],
		)
		self.assertEqual(validate_query({**query, "sortBy": "available"})["sortBy"], "available")
		with self.assertRaises(Exception):
			validate_query({**query, "sortBy": "image"})
		for legacy in ("quantity", "name", "batch"):
			with self.assertRaises(Exception):
				validate_query({**query, "sortBy": legacy})
		with self.assertRaises(Exception):
			validate_query(self.query("movements", "records", sortBy="stock_qty"))
		for dataset, view, fields in (
			(
				"movements",
				"lines",
				("document", "doctype", "borrower", "source_warehouse", "destination_warehouse"),
			),
			(
				"loans",
				"outcomes",
				(
					"document",
					"doctype",
					"kind",
					"docstatus",
					"borrower",
					"source_warehouse",
					"destination_warehouse",
				),
			),
		):
			for field in fields:
				self.assertEqual(validate_query(self.query(dataset, view, sortBy=field))["sortBy"], field)
		for changes in (
			{"grouping": "item"},
			{"groupBy": ["item", "item"]},
			{"groupBy": ["category", "location"]},
			{"groupBy": ["category", "item", "location"]},
		):
			with self.assertRaises(Exception):
				validate_query({**query, **changes})
		self.assertEqual(
			validate_query({**query, "groupBy": ["item", "category"]})["groupBy"],
			["item", "category"],
		)
		for dataset, view, invalid, _valid in (
			("movements", "summary", ["category"], ["category", "item"]),
			("loans", "items", ["item", "category"], ["category", "item"]),
		):
			with self.assertRaises(Exception):
				validate_query(self.query(dataset, view, groupBy=invalid))
		self.assertEqual(
			validate_query(self.query("loans", "items", groupBy=["item"]))["groupBy"],
			["item"],
		)

	def test_movement_groups_keep_kind_status_quantities_and_distinct_item_union(self):
		self.receipt(8)
		issue = self.movement("Issue", 3, self.a, None)
		self.submit(self.create(issue))
		query = self.query("movements", "lines", groupBy=["item"])
		page = report_page(query)
		self.assertEqual(page["matched_count"], 1)
		group = page["rows"][0]
		self.assertEqual(group["item_count"], 1)
		self.assertEqual(group["document_count"], 2)
		self.assertIsNone(group["stock_qty"])
		self.assertEqual(
			{(part["kind"], part["docstatus"]) for part in group["quantity_breakdown"]},
			{("Receive", 1), ("Issue", 1)},
		)
		by_kind = {part["kind"]: part for part in group["quantity_breakdown"]}
		self.assertEqual(by_kind["Receive"]["stock_qty"], 8)
		self.assertEqual(by_kind["Issue"]["stock_qty"], 3)
		self.assertEqual(by_kind["Receive"]["available_effect"], 8)
		self.assertEqual(by_kind["Issue"]["available_effect"], -3)
		self.assertEqual(page["summary"], report_page({**query, "groupBy": []})["summary"])

	def test_two_level_group_pages_include_full_authorized_groups_then_lazy_children(self):
		self.receipt(4)
		query = self.query("inventory", "summary", groupBy=["category", "item"])
		with self.assertRaises(Exception):
			report_page(query, table="group_source")
		page = report_page(query, page_length=1)
		self.assertEqual(page["table"], "group_roots")
		self.assertEqual(page["rows"][0]["group_path"], [""])
		self.assertEqual(page["rows"][0]["row_count"], 1)
		self.assertEqual(page["rows"][0]["allocation_count"], 0)
		roots = group_page(query, parent_path=[], page_length=1)
		self.assertEqual(roots["row_grain"], "group")
		self.assertEqual(roots["rows"][0]["group_path"], [""])
		children = group_page(
			query,
			parent_path=[""],
			parent_unit=roots["rows"][0]["stock_uom"],
			page_length=1,
		)
		self.assertEqual(children["rows"][0]["group_path"], ["", self.item])
		allocations = group_page(
			query,
			parent_path=["", self.item],
			parent_unit=children["rows"][0]["stock_uom"],
			page_length=1,
		)
		self.assertEqual(allocations["row_grain"], "stock")
		self.assertEqual(allocations["matched_count"], 1)

	def test_grouped_report_leaf_columns_are_curated_and_item_groups_include_authorized_images(self):
		self.receipt(4)
		image = "/files/grouped-item-preview.png"
		frappe.db.set_value("Item", self.item, "image", image)
		query = self.query("inventory", "summary", groupBy=["item"])
		page = report_page(query)
		self.assertNotIn("image", page["leaf_columns"])
		self.assertNotIn("batch_tracked", page["leaf_columns"])
		self.assertIn("category", page["leaf_columns"])
		self.assertEqual(page["rows"][0]["image"], image)
		groups = group_page(query, parent_path=[])
		self.assertEqual(groups["rows"][0]["item_code"], self.item)
		self.assertEqual(groups["rows"][0]["image"], image)

	def test_group_page_sorts_complete_filtered_groups_and_pages_matching_allocations(self):
		self.receipt(8)
		self.receipt(4)
		issue = self.movement("Issue", 3, self.a, None)
		self.submit(self.create(issue))
		query = self.query("movements", "lines", groupBy=["kind"], sortBy="stock_qty", sort="desc")
		page = report_page(query, page_length=1)
		self.assertEqual(page["matched_count"], 2)
		self.assertEqual(page["next_start"], 1)
		self.assertIn("document", page["leaf_columns"])
		self.assertNotIn("id", page["leaf_columns"])
		self.assertEqual(page["rows"][0]["kind"], "Receive")
		self.assertEqual(page["rows"][0]["stock_qty"], 12)
		groups = group_page(query, parent_path=[], page_length=1)
		self.assertEqual(groups["matched_count"], 2)
		self.assertEqual(groups["next_start"], 1)
		self.assertEqual(groups["rows"][0]["kind"], "Receive")
		children = group_page(
			query,
			parent_path=["Receive"],
			parent_unit=groups["rows"][0]["stock_uom"],
			page_length=1,
		)
		self.assertEqual(children["matched_count"], 2)
		self.assertEqual(children["next_start"], 1)
		self.assertEqual(children["row_grain"], "allocation")
		self.assertEqual(children["rows"][0]["stock_qty"], 8)

	def test_group_page_keeps_same_group_path_separate_across_native_units(self):
		self.receipt(5)
		uom_name = "TIUOM-" + self.token
		frappe.get_doc({"doctype": "UOM", "uom_name": uom_name}).insert()
		other = frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": "TI-UNIT-" + self.token,
				"item_name": "Native test item in alternate unit",
				"stock_uom": uom_name,
				"item_group": "All Item Groups",
				"is_stock_item": 1,
			}
		).insert()
		data = self.movement(quantity=7)
		data["items"][0].update({"item_code": other.name, "uom": uom_name, "conversion_factor": 1})
		self.submit(self.create(data))
		query = self.query("inventory", "summary", groupBy=["category", "item"])
		query["scope"]["item"] = ""
		roots = group_page(query, parent_path=[], page_length=1)
		self.assertEqual(roots["matched_count"], 2)
		self.assertEqual(roots["next_start"], 1)
		second_root = group_page(query, parent_path=[], start=1, page_length=1)["rows"][0]
		self.assertNotEqual(roots["rows"][0]["stock_uom"], second_root["stock_uom"])
		for root in (roots["rows"][0], second_root):
			children = group_page(
				query,
				parent_path=root["group_path"],
				parent_unit=root["stock_uom"],
				page_length=10,
			)
			self.assertEqual(children["matched_count"], 1)
			leaf = children["rows"][0]
			allocations = group_page(
				query,
				parent_path=leaf["group_path"],
				parent_unit=leaf["stock_uom"],
				page_length=10,
			)
			self.assertEqual(allocations["matched_count"], 1)
			self.assertEqual(allocations["rows"][0]["item_code"], leaf["item_code"])
		item_class = type(frappe.get_doc("Item", other.name))
		check_permission = item_class.check_permission

		def denied_item(item_doc, *args, **kwargs):
			if item_doc.name == other.name:
				raise frappe.PermissionError("Alternate unit Item denied")
			return check_permission(item_doc, *args, **kwargs)

		with patch.object(item_class, "check_permission", denied_item):
			restricted = group_page(query, parent_path=[], page_length=10)
			self.assertEqual(restricted["matched_count"], 1)
			self.assertNotEqual(restricted["rows"][0]["stock_uom"], uom_name)

	def test_grouped_native_records_keep_full_document_quantities_after_line_matching(self):
		other = frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": "TI-RECORD-" + self.token,
				"item_name": "Second native record item",
				"stock_uom": "Nos",
				"item_group": "All Item Groups",
				"is_stock_item": 1,
			}
		).insert()
		data = self.movement(quantity=5)
		data["items"].append({**data["items"][0], "item_code": other.name, "qty": 7})
		doc = self.submit(self.create(data))
		record_query = self.query("movements", "records", record=doc.name, statuses=[1])
		records = report_page(record_query)
		self.assertEqual(records["matched_count"], 1)
		self.assertEqual(records["rows"][0]["stock_qty"], 12)
		self.assertEqual(records["rows"][0]["allocation_count"], 2)
		self.assertEqual(records["rows"][0]["item_count"], 2)
		group_query = {**record_query, "groupBy": ["kind"]}
		grouped = report_page(group_query)
		self.assertEqual(grouped["rows"][0]["stock_qty"], 12)
		self.assertEqual(grouped["rows"][0]["document_count"], 1)
		child_records = group_page(group_query, parent_path=["Receive"], parent_unit="Nos")
		self.assertEqual(child_records["row_grain"], "record")
		self.assertEqual(child_records["matched_count"], 1)
		self.assertEqual(child_records["rows"][0]["stock_qty"], 12)
		matching_lines = report_page(self.query("movements", "lines", record=doc.name, statuses=[1]))
		self.assertEqual(matching_lines["matched_count"], 1)
		self.assertEqual(matching_lines["rows"][0]["stock_qty"], 5)
		self.submit(self.create(self.movement(quantity=2)))
		overlapping = report_page(self.query("movements", "records", statuses=[1], groupBy=["kind"]))
		self.assertEqual(overlapping["rows"][0]["document_count"], 2)
		self.assertEqual(overlapping["rows"][0]["item_count"], 2)
		self.assertEqual(overlapping["rows"][0]["stock_qty"], 14)

	def test_quarter_periods_use_site_date_and_calendar_boundaries(self):
		from temple_inventory.periods import resolve_period

		def resolve(key, today):
			return resolve_period({"period_key": key}, today_provider=lambda: today)

		self.assertEqual(
			resolve("this_quarter", "2026-04-05"),
			{"key": "this_quarter", "date_from": "2026-04-01", "date_to": "2026-04-05"},
		)
		self.assertEqual(
			resolve("last_quarter", "2026-01-19"),
			{"key": "last_quarter", "date_from": "2025-10-01", "date_to": "2025-12-31"},
		)

	def test_read_cache_round_trips_base_sqlite_projection_into_private_database(self):
		from types import SimpleNamespace

		from frappe.utils import now_datetime

		from temple_inventory import read_cache
		from temple_inventory.report_query import resolved_period

		class Cache:
			def __init__(self):
				self.values = {}

			def make_key(self, key):
				return key

			def get(self, key):
				return self.values.get(key)

			def set(self, key, value, ex=None):
				self.values[key] = value

		cache = Cache()
		quarter = resolved_period({"key": "this_quarter", "from": "", "to": ""})
		base_identity = SimpleNamespace(
			query={
				"dataset": "inventory",
				"view": "summary",
				"groupBy": ["category"],
				"sortBy": "available",
				"sort": "asc",
				"statuses": [1, 0],
				"itemCodes": ["B", "A"],
				"search": " Tea ",
				"scope": {
					"warehouses": ["Z", "A"],
					"period": {"key": "this_quarter", "from": "", "to": ""},
				},
			},
		)
		reordered_identity = SimpleNamespace(
			query={
				"dataset": "inventory",
				"view": "summary",
				"groupBy": ["item"],
				"sortBy": "item_name",
				"sort": "desc",
				"statuses": [0, 1],
				"itemCodes": ["A", "B"],
				"search": "tea",
				"scope": {
					"warehouses": ["A", "Z"],
					"period": {
						"key": "custom",
						"from": quarter["date_from"],
						"to": quarter["date_to"],
					},
				},
			},
		)
		identity_args = ("generation", "permission", "configuration")
		self.assertEqual(
			read_cache._projection_identity(base_identity, *identity_args),
			read_cache._projection_identity(reordered_identity, *identity_args),
		)
		with (
			patch.object(read_cache, "_request_cacheable", return_value=True),
			patch.object(read_cache, "_redis", return_value=cache),
			patch.object(read_cache, "data_generation", return_value="4"),
			patch.object(read_cache, "_permission_signature", return_value="account-grants"),
			patch.object(read_cache, "_configuration_signature", return_value="settings"),
		):
			stamp = str(now_datetime())
			with tempfile.TemporaryDirectory() as base:
				first = SimpleNamespace(
					query={"dataset": "inventory", "view": "summary", "groupBy": ["category"]},
					cache_bypass=False,
					cache_context={
						"generation": "4",
						"permission": "account-grants",
						"configuration": "settings",
						"user": frappe.session.user,
					},
					temp=SimpleNamespace(name=base),
					context=SimpleNamespace(snapshot=stamp),
					column_names={"summary": ["id", "available"]},
					main_table="summary",
					summary_table="summary",
					sums=("available",),
				)
				first.db = sqlite3.connect(f"{base}/projection.sqlite")
				first.db.execute("create table rows (table_id text, identity text, payload text)")
				first.db.execute("insert into rows values ('summary', 'A', '{\"available\":3}')")
				read_cache.store_projection(first)
				first.db.close()

				second_dir = tempfile.TemporaryDirectory()
				second_path = f"{second_dir.name}/projection.sqlite"
				second = SimpleNamespace(
					query={"dataset": "inventory", "view": "summary", "groupBy": ["item"]},
					cache_bypass=False,
					cache_context=first.cache_context,
					temp=second_dir,
					context=SimpleNamespace(snapshot=str(now_datetime())),
					db=sqlite3.connect(second_path),
				)
				second.db.execute("create table rows (table_id text, identity text, payload text)")
				self.assertTrue(read_cache.restore_projection(second))
				self.assertEqual(second.main_table, "summary")
				self.assertEqual(second.sums, ("available",))
				self.assertEqual(second.db.execute("select identity from rows").fetchone()[0], "A")
				self.assertEqual(second.db.execute("pragma database_list").fetchone()[2], second_path)
				second.db.close()
				second_dir.cleanup()

				# Reject snapshots after either the shared source generation or the
				# current account's grants have changed, even when the key still exists.
				for changed in (
					{"data_generation": "5"},
					{"_permission_signature": "revoked-grant"},
				):
					rejected_dir = tempfile.TemporaryDirectory()
					rejected_path = f"{rejected_dir.name}/projection.sqlite"
					rejected = SimpleNamespace(
						query={"dataset": "inventory", "view": "summary", "groupBy": ["item"]},
						cache_bypass=False,
						cache_context=first.cache_context,
						temp=rejected_dir,
						context=SimpleNamespace(snapshot=str(now_datetime())),
						db=sqlite3.connect(rejected_path),
					)
					rejected.db.execute("create table rows (table_id text, identity text, payload text)")
					with (
						patch.object(read_cache, "_request_cacheable", return_value=True),
						patch.multiple(
							read_cache,
							**changed,
						),
					):
						self.assertFalse(read_cache.restore_projection(rejected))
					rejected.db.close()
					rejected_dir.cleanup()

				cache.values.clear()
				with patch.object(read_cache, "MAX_ENTRY_BYTES", 1):
					first.db = sqlite3.connect(f"{base}/projection.sqlite")
					read_cache.store_projection(first)
					first.db.close()
				self.assertEqual(cache.values, {})

				# A disconnected Redis must turn reuse/store into a cache miss, not a
				# projection error.
				import redis

				with patch.object(read_cache, "_redis", side_effect=redis.exceptions.ConnectionError):
					self.assertFalse(read_cache.restore_projection(second))
					first.db = sqlite3.connect(f"{base}/projection.sqlite")
					read_cache.store_projection(first)
					first.db.close()

	def test_read_cache_snapshot_age_uses_calculation_start_and_expires_at_ttl(self):
		from datetime import timedelta

		from temple_inventory import read_cache

		now = frappe.utils.now_datetime()
		with patch.object(read_cache, "now_datetime", return_value=now):
			self.assertTrue(read_cache._fresh(str(now - timedelta(seconds=14))))
			self.assertFalse(read_cache._fresh(str(now - timedelta(seconds=15))))
			self.assertFalse(read_cache._fresh(str(now + timedelta(seconds=1))))

	def test_durable_cache_revision_update_rolls_back_with_its_source_transaction(self):
		from temple_inventory import read_cache

		def get_revision():
			return frappe.db.sql(
				"select value from tabSingles where doctype=%s and field='ti_read_generation'",
				("Temple Inventory Settings",),
				as_dict=True,
			)

		before_rows = get_revision()
		before = int(before_rows[0].value) if before_rows else 0
		frappe.db.savepoint("ti_cache_revision")
		read_cache._bump_durable_revision()
		after = int(get_revision()[0].value)
		self.assertEqual(after, before + 1)
		frappe.db.rollback(save_point="ti_cache_revision")
		after_rollback_rows = get_revision()
		after_rollback = int(after_rollback_rows[0].value) if after_rollback_rows else 0
		self.assertEqual(after_rollback, before)

	def test_source_document_event_advances_revision_inside_its_transaction(self):
		from temple_inventory import read_cache

		def get_revision():
			rows = frappe.db.sql(
				"select value from tabSingles where doctype=%s and field='ti_read_generation'",
				("Temple Inventory Settings",),
				as_dict=True,
			)
			return int(rows[0].value) if rows else 0

		before = get_revision()
		frappe.db.savepoint("ti_source_cache_hook")
		after_commit = frappe.db.after_commit
		after_rollback = frappe.db.after_rollback
		commit_callbacks = list(after_commit._functions)
		rollback_callbacks = list(after_rollback._functions)
		old_scheduled = getattr(frappe.local, "inventory_cache_dirty_scheduled", False)
		old_touched = getattr(frappe.local, "inventory_cache_revision_touched", False)
		try:
			frappe.local.inventory_cache_dirty_scheduled = False
			frappe.local.inventory_cache_revision_touched = False
			read_cache.invalidate_after_commit(frappe._dict(doctype="Item"), "on_update")
			self.assertEqual(get_revision(), before + 1)
			self.assertEqual(len(after_commit._functions), len(commit_callbacks) + 1)
			frappe.db.rollback(save_point="ti_source_cache_hook")
			self.assertEqual(get_revision(), before)
		finally:
			after_commit._functions.clear()
			after_commit._functions.extend(commit_callbacks)
			after_rollback._functions.clear()
			after_rollback._functions.extend(rollback_callbacks)
			frappe.local.inventory_cache_dirty_scheduled = old_scheduled
			frappe.local.inventory_cache_revision_touched = old_touched

	def test_read_cache_redis_outage_falls_back_to_fresh_projection(self):
		import redis

		from temple_inventory import read_cache

		self.receipt(2)
		frappe.local.inventory_cache_request_generation = "live-generation"
		frappe.local.inventory_cache_request_started = str(frappe.utils.now_datetime())
		with (
			patch.object(read_cache, "_request_cacheable", return_value=True),
			patch.object(read_cache, "_redis", side_effect=redis.exceptions.ConnectionError),
		):
			result = report_page(self.query("inventory", "summary"), refresh=False)
		self.assertGreater(result["matched_count"], 0)

	def test_explicit_refresh_and_test_or_write_contexts_bypass_shared_cache(self):
		import redis

		from temple_inventory import read_cache
		from temple_inventory.projections import Projection

		self.receipt(2)
		query = self.query("inventory", "summary")
		frappe.local.inventory_cache_request_generation = "live-generation"
		frappe.local.inventory_cache_request_started = str(frappe.utils.now_datetime())
		self.assertFalse(read_cache._request_cacheable())
		with patch.object(read_cache, "_redis", side_effect=redis.exceptions.ConnectionError):
			result = report_page(query, refresh=True)
		self.assertGreater(result["matched_count"], 0)
		self.assertGreater(frappe.db.transaction_writes, 0)
		with Projection(query, refresh=True) as fresh:
			self.assertNotEqual(fresh.calculation_started, "2000-01-01 00:00:00")

	def test_user_role_edit_and_global_clear_cache_advance_durable_revision(self):
		from temple_inventory import read_cache

		after_commit = frappe.db.after_commit
		after_rollback = frappe.db.after_rollback
		prior_commit_callbacks = list(after_commit._functions)
		prior_rollback_callbacks = list(after_rollback._functions)
		old_scheduled = getattr(frappe.local, "inventory_cache_dirty_scheduled", False)
		old_touched = getattr(frappe.local, "inventory_cache_revision_touched", False)
		with (
			patch.object(read_cache, "_bump_durable_revision") as bump,
			patch.object(read_cache, "invalidate_now") as invalidate,
		):
			frappe.local.inventory_cache_dirty_scheduled = False
			frappe.local.inventory_cache_revision_touched = False
			before = frappe._dict(roles=[frappe._dict(role="Stock User")])
			user = frappe._dict(
				doctype="User",
				roles=[frappe._dict(role="System Manager")],
				get_doc_before_save=lambda: before,
				has_value_changed=lambda _field: False,
			)
			read_cache.invalidate_after_commit(user, "on_update")
			self.assertEqual(bump.call_count, 1)
			rollback_callback = list(after_rollback._functions)[-1]
			rollback_callback()
			self.assertEqual(invalidate.call_count, 0)

			frappe.local.inventory_cache_dirty_scheduled = False
			frappe.local.inventory_cache_revision_touched = False
			read_cache.invalidate_after_clear_cache()
			self.assertEqual(bump.call_count, 2)
			commit_callback = list(after_commit._functions)[-1]
			commit_callback()
			self.assertEqual(invalidate.call_count, 1)
		after_commit._functions.clear()
		after_commit._functions.extend(prior_commit_callbacks)
		after_rollback._functions.clear()
		after_rollback._functions.extend(prior_rollback_callbacks)
		frappe.local.inventory_cache_dirty_scheduled = old_scheduled
		frappe.local.inventory_cache_revision_touched = old_touched

	def test_permission_signature_tracks_persisted_role_revocation(self):
		from frappe.cache_manager import clear_defaults_cache

		from temple_inventory import read_cache

		user = frappe.get_doc(
			{
				"doctype": "User",
				"email": f"ti-cache-role-{self.token}@example.test",
				"first_name": "Cache permission owner",
				"send_welcome_email": 0,
				"user_type": "System User",
				"roles": [{"role": "System Manager"}, {"role": "Stock User"}],
			}
		).insert()
		self.addCleanup(clear_defaults_cache, user.name)
		before = read_cache._permission_signature(user.name)
		user.set("roles", [{"role": "Stock User"}])
		user.save()
		after = read_cache._permission_signature(user.name)
		self.assertNotEqual(before, after)
		# Simulate a warm Frappe role list from before the revoke. The app gate
		# must still deny the next read using the current database grants.
		from temple_inventory.native_access import require_frontend_account

		previous_user = frappe.session.user
		try:
			frappe.set_user(user.name)
			with patch("frappe.get_roles", return_value=["System Manager", "Stock User"]):
				with self.assertRaises(frappe.PermissionError):
					require_frontend_account()
		finally:
			frappe.set_user(previous_user)

	def test_read_cache_invalidation_publishes_after_commit_and_ignores_rollback(self):
		from temple_inventory import read_cache

		after_commit = frappe.db.after_commit
		after_rollback = frappe.db.after_rollback
		prior_commit_callbacks = list(after_commit._functions)
		prior_rollback_callbacks = list(after_rollback._functions)
		old_scheduled = getattr(frappe.local, "inventory_cache_dirty_scheduled", False)
		old_touched = getattr(frappe.local, "inventory_cache_revision_touched", False)
		with (
			patch.object(read_cache, "invalidate_now") as invalidate,
		):
			try:
				after_commit._functions.clear()
				after_rollback._functions.clear()
				frappe.local.inventory_cache_dirty_scheduled = False
				frappe.local.inventory_cache_revision_touched = False
				read_cache.invalidate_after_commit()
				self.assertEqual(invalidate.call_count, 0)
				rollback_callback = after_rollback._functions[-1]
				rollback_callback()
				self.assertEqual(invalidate.call_count, 0)
				read_cache.invalidate_after_commit()
				after_commit._functions[-1]()
				self.assertEqual(invalidate.call_count, 1)
				callback_count = len(after_commit._functions)
				read_cache.invalidate_after_commit(
					frappe._dict(
						doctype="User",
						get_doc_before_save=lambda: None,
						has_value_changed=lambda _field: False,
					),
					"on_update",
				)
				self.assertEqual(len(after_commit._functions), callback_count)
			finally:
				after_commit._functions.clear()
				after_commit._functions.extend(prior_commit_callbacks)
				after_rollback._functions.clear()
				after_rollback._functions.extend(prior_rollback_callbacks)
				frappe.local.inventory_cache_dirty_scheduled = old_scheduled
				frappe.local.inventory_cache_revision_touched = old_touched

	def test_read_cache_boundary_only_commits_clean_read_request(self):
		from temple_inventory import read_cache

		old_form_dict = getattr(frappe.local, "form_dict", None)
		frappe.local.form_dict = frappe._dict(cmd="temple_inventory.projections.report_page")
		with (
			patch.object(frappe.db, "transaction_writes", 0),
			patch.object(frappe.db, "before_commit", frappe._dict(_functions=[])),
			patch.object(frappe.db, "after_commit", frappe._dict(_functions=[])),
			patch.object(frappe.db, "before_rollback", frappe._dict(_functions=[])),
			patch.object(frappe.db, "after_rollback", frappe._dict(_functions=[])),
		):
			with (
				patch.object(frappe.db, "commit") as commit,
				patch.object(read_cache, "data_generation", return_value="17"),
			):
				frappe.db.transaction_writes = 1
				read_cache.begin_read_request()
				self.assertEqual(commit.call_count, 0)
				self.assertIsNone(frappe.local.inventory_cache_request_generation)
				frappe.db.transaction_writes = 0
				read_cache.begin_read_request()
				commit.assert_called_once_with()
				self.assertEqual(frappe.local.inventory_cache_request_generation, "17")
		frappe.local.form_dict = old_form_dict

	def test_browse_http_contract_reuses_redis_base_projection_across_groups(self):
		from temple_inventory import read_cache
		from temple_inventory.projections import stock_rows

		# This request intentionally exercises the real Redis path while the test's
		# database transaction contains uncommitted fixtures. Evict its snapshot
		# after teardown so a later live request cannot observe rolled-back rows.
		self.addCleanup(read_cache.invalidate_now)
		old_generation = getattr(frappe.local, "inventory_cache_request_generation", None)
		old_started = getattr(frappe.local, "inventory_cache_request_started", None)
		self.addCleanup(lambda: setattr(frappe.local, "inventory_cache_request_generation", old_generation))
		self.addCleanup(lambda: setattr(frappe.local, "inventory_cache_request_started", old_started))
		with patch.object(
			frappe.db, "commit", side_effect=AssertionError("cacheable projection committed test data")
		):
			self.receipt(3)
			query = self.query("inventory", "summary")
			query["scope"]["warehouses"] = [self.a, self.b]
			frappe.local.inventory_cache_request_generation = read_cache.data_generation()
			frappe.local.inventory_cache_request_started = str(frappe.utils.now_datetime())
			with (
				patch.object(read_cache, "_request_cacheable", return_value=True),
				patch("temple_inventory.projections.stock_rows", wraps=stock_rows) as source_reads,
			):
				first = report_page(query, page_length=1)
				second = report_page({**query, "groupBy": ["item"]}, page_length=1)
				third = report_page(
					{
						**query,
						"scope": {**query["scope"], "warehouses": [self.b, self.a]},
						"sortBy": "available",
						"sort": "desc",
					},
					start=1,
					page_length=1,
				)

			class LockTimeout:
				def acquire(self, blocking=True):
					return False

				def release(self):
					raise AssertionError("a lock that was not acquired must not be released")

			class WaiterCache:
				def make_key(self, key):
					return frappe.cache.make_key(key)

				def get(self, key):
					return frappe.cache.get(key)

				def set(self, key, value, ex=None):
					return frappe.cache.set(key, value, ex=ex)

				def lock(self, *_args, **_kwargs):
					return LockTimeout()

			with patch.object(read_cache, "_redis", return_value=WaiterCache()):
				waiter = report_page(query, start=1, page_length=1)
		self.assertGreater(first["matched_count"], 0)
		self.assertEqual(second["matched_count"], first["matched_count"])
		self.assertEqual(third["matched_count"], first["matched_count"])
		self.assertEqual(waiter["matched_count"], first["matched_count"])
		self.assertEqual(source_reads.call_count, 1)

	def test_redis_failure_after_cached_restore_rebuilds_from_fresh_source(self):
		from temple_inventory import read_cache
		from temple_inventory.projections import stock_rows

		self.addCleanup(read_cache.invalidate_now)
		old_request = getattr(frappe.local, "request", None)
		old_generation = getattr(frappe.local, "inventory_cache_request_generation", None)
		old_started = getattr(frappe.local, "inventory_cache_request_started", None)
		self.addCleanup(setattr, frappe.local, "inventory_cache_request_generation", old_generation)
		self.addCleanup(setattr, frappe.local, "inventory_cache_request_started", old_started)
		if old_request is None:
			self.addCleanup(lambda: setattr(frappe.local, "request", None))
		else:
			self.addCleanup(setattr, frappe.local, "request", old_request)
		self.receipt(3)
		frappe.local.request = frappe._dict(headers={"X-Inventory-Account": frappe.session.user})
		frappe.local.inventory_cache_request_generation = read_cache.data_generation()
		frappe.local.inventory_cache_request_started = str(frappe.utils.now_datetime())
		query = self.query("inventory", "summary")
		query["scope"]["warehouses"] = [self.a]
		with patch.object(read_cache, "_request_cacheable", return_value=True):
			first = report_page(query)
			self.assertGreater(first["matched_count"], 0)
			with (
				patch.object(read_cache, "cache_context_current", return_value=None),
				patch.object(read_cache, "restart_owned_read_snapshot", return_value=True) as restart,
				patch("temple_inventory.projections.stock_rows", wraps=stock_rows) as source_reads,
			):
				second = report_page(query)
		self.assertEqual(second["matched_count"], first["matched_count"])
		restart.assert_called_once_with()
		source_reads.assert_called_once()

	def test_expired_cached_projection_after_grouping_rebuilds_before_serving(self):
		from temple_inventory import read_cache
		from temple_inventory.projections import stock_rows

		self.addCleanup(read_cache.invalidate_now)
		old_request = getattr(frappe.local, "request", None)
		old_generation = getattr(frappe.local, "inventory_cache_request_generation", None)
		old_started = getattr(frappe.local, "inventory_cache_request_started", None)
		self.addCleanup(setattr, frappe.local, "inventory_cache_request_generation", old_generation)
		self.addCleanup(setattr, frappe.local, "inventory_cache_request_started", old_started)
		if old_request is None:
			self.addCleanup(lambda: setattr(frappe.local, "request", None))
		else:
			self.addCleanup(setattr, frappe.local, "request", old_request)
		self.receipt(3)
		frappe.local.request = frappe._dict(headers={"X-Inventory-Account": frappe.session.user})
		frappe.local.inventory_cache_request_generation = read_cache.data_generation()
		frappe.local.inventory_cache_request_started = str(frappe.utils.now_datetime())
		query = self.query("inventory", "summary")
		query["groupBy"] = ["category"]
		query["scope"]["warehouses"] = [self.a]
		with patch.object(read_cache, "_request_cacheable", return_value=True):
			first = group_page(query, parent_path=[])
			self.assertGreater(first["matched_count"], 0)
			original_stock_rows = stock_rows

			def fresh_rows(context):
				for row in original_stock_rows(context):
					row["category"] = "fresh-source-marker"
					yield row

			restore_results = []
			original_restore = read_cache.restore_projection

			def tracked_restore(projection):
				result = original_restore(projection)
				restore_results.append(result)
				return result

			with (
				patch.object(read_cache, "_fresh", side_effect=[True, False]),
				patch.object(read_cache, "restore_projection", side_effect=tracked_restore),
				patch.object(read_cache, "restart_owned_read_snapshot", return_value=True) as restart,
				patch("temple_inventory.projections.stock_rows", side_effect=fresh_rows) as source_reads,
			):
				second = group_page(query, parent_path=[])

		self.assertEqual(restore_results, [True, False])
		restart.assert_called_once_with()
		self.assertEqual(source_reads.call_count, 1)
		self.assertTrue(any("fresh-source-marker" in (row.get("group_path") or []) for row in second["rows"]))

	def test_uncached_snapshot_restart_rejects_pending_writes(self):
		from temple_inventory import read_cache

		with (
			patch.object(read_cache, "_request_cacheable", return_value=True),
			patch.object(frappe.db, "transaction_writes", 1),
		):
			self.assertFalse(read_cache.restart_owned_read_snapshot())

	def test_inflight_projection_discards_build_after_generation_change(self):
		from temple_inventory import read_cache
		from temple_inventory.projections import stock_rows

		self.addCleanup(read_cache.invalidate_now)
		old_generation = getattr(frappe.local, "inventory_cache_request_generation", None)
		old_started = getattr(frappe.local, "inventory_cache_request_started", None)
		self.addCleanup(setattr, frappe.local, "inventory_cache_request_generation", old_generation)
		self.addCleanup(setattr, frappe.local, "inventory_cache_request_started", old_started)
		frappe.local.inventory_cache_request_generation = read_cache.data_generation()
		frappe.local.inventory_cache_request_started = str(frappe.utils.now_datetime())
		original = stock_rows

		def concurrent_generation_change(context):
			rows = original(context)
			read_cache.invalidate_now()
			return rows

		query = self.query("inventory", "summary")
		query["scope"]["warehouses"] = [self.a]
		with (
			patch.object(read_cache, "_request_cacheable", return_value=True),
			patch("temple_inventory.projections.stock_rows", side_effect=concurrent_generation_change),
			self.assertRaises(frappe.ValidationError),
		):
			report_page(query)

	def test_saved_record_scope_requires_exact_doctype_name_and_status(self):
		self.receipt(10)
		for doctype, dataset in (("Stock Entry", "movements"), ("Stock Reconciliation", "counts")):
			draft = (
				self.create(self.movement())
				if doctype == "Stock Entry"
				else counts.save_count(str(uuid4()), self.count_data(9, "Reviewed"))
			)
			for status in (0, 1, 2):
				if status == 1:
					draft = native_api.submit_document(doctype, draft.name, draft.version)
				elif status == 2:
					draft = native_api.cancel_document(doctype, draft.name, draft.version)
				query = self.query(
					dataset, "records", record=draft.name, recordDoctype=doctype, statuses=[status]
				)
				result = report_page(query)
				self.assertGreater(result["matched_count"], 0)
				self.assertEqual({row["document"] for row in result["rows"]}, {draft.name})
				other = "Stock Reconciliation" if doctype == "Stock Entry" else "Stock Entry"
				self.assertEqual(report_page({**query, "recordDoctype": other})["matched_count"], 0)
				with Projection(query) as projection:
					projection.build(related=True)
					self.assertEqual({row["document"] for row in projection.rows("lines")}, {draft.name})
					self.assertIn("line_name", projection.column_names["lines"])
					self.assertIn("document_id", projection.column_names["lines"])
					self.assertIn("id", projection.column_names["lines"])

	def test_saved_opening_export_uses_native_movement_quantities_and_excludes_counts(self):
		doc = frappe.get_doc(
			{
				"doctype": "Stock Reconciliation",
				"company": self.company,
				"purpose": "Opening Stock",
				"expense_account": frappe.db.get_value(
					"Account", {"company": self.company, "account_name": "Temporary Opening"}, "name"
				),
				"posting_date": nowdate(),
				"items": [{"item_code": self.item, "warehouse": self.a, "qty": 7, "valuation_rate": 10}],
			}
		).insert()
		for status in (0, 1, 2):
			if status == 1:
				doc.submit()
			elif status == 2:
				doc.cancel()
			loaded = native_api.load_document(doc.doctype, doc.name)
			self.assertEqual(loaded.purpose, "Opening Stock")
			query = self.query(
				"movements", "records", record=doc.name, recordDoctype=doc.doctype, statuses=[status]
			)
			page = report_page(query, page_length=5)
			self.assertEqual(page["matched_count"], 1)
			self.assertEqual(page["rows"][0]["document"], doc.name)
			self.assertEqual(page["rows"][0]["kind"], "Opening")
			with Projection(query) as projection:
				projection.build(related=True)
				lines = list(projection.rows("lines"))
				self.assertEqual(len(lines), 1)
				self.assertEqual(lines[0]["stock_qty"], 7)
				self.assertEqual(lines[0]["actual_qty"], 7)
				self.assertEqual(lines[0]["doctype"], doc.doctype)
				self.assertEqual(lines[0]["line_name"], doc.items[0].name)
			self.assertEqual(report_page({**query, "dataset": "counts"})["matched_count"], 0)
			self.assertEqual(report_page({**query, "recordDoctype": "Stock Entry"})["matched_count"], 0)

	def test_empty_restrictive_item_and_allocation_selections_never_widen_scope(self):
		self.receipt()
		for dataset in DEFINITIONS:
			for flag, values in (("restrictItems", "itemCodes"), ("restrictAllocations", "allocationKeys")):
				result = report_page(self.query(dataset, **{flag: True, values: []}))
				self.assertEqual(result["matched_count"], 0, (dataset, flag))
		self.assertEqual(
			report_page(self.query(scope={"warehouses": ["missing-warehouse"]}, includeZero=True))[
				"matched_count"
			],
			0,
		)

	def test_native_pagination_and_complete_unit_summary_are_independent(self):
		data = self.movement(quantity=1)
		data["items"] = [copy.deepcopy(data["items"][0]) for _ in range(205)]
		entry = self.submit(self.create(data))
		query = self.query("movements", "lines", record=entry.name)
		first = report_page(query, page_length=30)
		self.assertEqual(len(first["rows"]), 30)
		self.assertEqual(first["matched_count"], 205)
		self.assertEqual(first["next_start"], 30)
		self.assertEqual(sum(row["stock_qty"] for row in first["summary"]), 205)
		last = report_page(query, start=200, page_length=30)
		self.assertEqual(len(last["rows"]), 5)
		self.assertIsNone(last["next_start"])
		self.assertFalse({row["id"] for row in first["rows"]} & {row["id"] for row in last["rows"]})

	def test_current_stock_does_not_ignore_transaction_only_context(self):
		entry = self.receipt(5)
		self.assertGreater(report_page(self.query())["matched_count"], 0)
		for scope in (
			{"kinds": ["Receive"]},
			{"sources": [self.a]},
			{"destinations": [self.a]},
			{"activity": self.activity.name},
		):
			query = self.query(scope={"item": self.item, "warehouses": [self.root], **scope})
			self.assertEqual(report_page(query)["matched_count"], 0, scope)
		for restriction in ({"record": entry.name}, {"loan": entry.name}):
			self.assertEqual(report_page(self.query(**restriction))["matched_count"], 0)

	def test_historical_stock_has_both_transfer_branches_and_one_reconciliation_adjustment(self):
		self.receipt(10)
		self.submit(self.create(self.movement("Transfer", 3, self.a, self.b)))
		doc = counts.save_count(str(uuid4()), self.count_data(5, "Reviewed"))
		native_api.submit_document("Stock Reconciliation", doc.name, doc.version)
		query = self.query("inventory", "historical")
		query["scope"]["period"] = {"key": "custom", "from": nowdate(), "to": nowdate()}
		result = report_page(query)
		rows = {row["warehouse"]: row for row in result["rows"]}
		self.assertEqual(rows[self.a]["closing"], 5)
		self.assertEqual(rows[self.b]["closing"], 3)
		self.assertEqual(sum(row["adjustment"] for row in rows.values()), -2)
		for row in rows.values():
			self.assertEqual(
				row["opening"] + row["incoming"] - row["outgoing"] + row["adjustment"], row["closing"]
			)
		ledger = report_page(query, table="ledger")
		self.assertEqual(len([row for row in ledger["rows"] if row["document"] == doc.name]), 1)
		query["scope"]["warehouses"] = [self.a]
		self.assertEqual(report_page(query)["rows"][0]["closing"], 5)

	def test_grouped_projection_retains_exact_allocation_detail_and_original_loan_lines(self):
		loan = self.loan(5)
		self.submit(self.settlement(loan, 2))
		result = report_page(self.query("loans", "outcomes", groupBy=["item"]))
		self.assertEqual(result["rows"][0]["outstanding"], 3)
		self.assertIn("outcomes", {table["id"] for table in result["tables"]})
		settlements = report_page(self.query("loans", "settlements"))
		self.assertEqual(settlements["rows"][0]["original_loan_line"], loan["items"][0].name)

	def test_activity_indirect_settlement_relation_follows_exact_original_line(self):
		loan = self.loan()
		settlement = self.settlement(loan)
		# A compatible native Desk settlement may retain the original link without
		# repeating its activity. The native model permits this only when blank.
		desk = frappe.get_doc("Stock Entry", settlement.name)
		desk.ti_activity = None
		desk.save()
		desk.submit()
		query = self.query("activities", "records")
		query["scope"]["activity"] = self.activity.name
		result = report_page(query)
		linked = [row for row in result["rows"] if row["document"] == settlement.name]
		self.assertEqual(len(linked), 1)
		self.assertEqual(linked[0]["relationship"], "original_loan")
		metadata = report_page({**query, "view": "metadata"})["rows"][0]
		self.assertEqual(metadata["count_Loan"], 1)
		self.assertEqual(metadata["count_Return"], 0)
		overview = activities.load_activity(self.activity.name)
		self.assertEqual(overview["summary"]["count_Loan"], 1)
		self.assertEqual(overview["summary"]["count_Return"], 0)
		self.assertEqual(overview["outstanding_line_count"], 1)
		self.assertEqual(overview["loan_totals"][0]["outstanding"], 3)
		self.assertEqual(overview["movement_totals"][0]["kind"], "Loan")
		self.assertTrue(overview["can_write"])

	def test_activity_uncategorized_scope_requires_a_direct_root_assignment(self):
		category = catalog.create_category("Activity category " + self.token)
		frappe.db.set_value("Item", self.item, "item_group", category["name"])
		self.loan(2)
		query = self.query("activities", "metadata")
		query["scope"]["uncategorized"] = True
		self.assertEqual(report_page(query)["matched_count"], 0)

		frappe.db.set_value("Item", self.item, "item_group", catalog.item_group_root())
		self.assertEqual(report_page(query)["matched_count"], 1)

	def test_counts_report_contains_only_actual_native_reconciliations(self):
		self.receipt(10)
		counts.save_count(str(uuid4()), self.count_data(10, "Reviewed"))
		self.assertEqual(report_page(self.query("counts"))["matched_count"], 0)
		doc = counts.save_count(str(uuid4()), self.count_data(9, "Reviewed"))
		native_api.submit_document("Stock Reconciliation", doc.name, doc.version)
		rows = report_page(self.query("counts", "lines"))["rows"]
		self.assertEqual(rows[0]["difference"], -1)
		self.assertEqual(rows[0]["review_state"], "Reviewed")

	def test_named_views_preserve_relative_dates_and_account_isolation(self):
		from frappe.cache_manager import clear_defaults_cache

		self.addCleanup(clear_defaults_cache, "Administrator")
		query = self.query("movements")
		query["scope"]["period"]["key"] = "last_30_days"
		view = preferences.save_named_view("常用 " + self.token, query)
		self.assertEqual(view["query"]["scope"]["period"]["key"], "last_30_days")
		self.assertEqual(view["query"]["scope"]["period"]["from"], "")
		self.assertIn(view["id"], {value["id"] for value in preferences.named_views()})
		preferences.save_named_view("改名 " + self.token, view_id=view["id"], revision=view["revision"])
		with self.assertRaises(frappe.TimestampMismatchError):
			preferences.delete_named_view(view["id"], view["revision"])

	def test_xlsx_and_csv_keep_identical_selected_numeric_table_scope(self):
		self.receipt(12)
		query = self.query("movements", "lines")
		with tempfile.TemporaryDirectory() as directory, Projection(query) as projection:
			projection.build(related=True)
			request = frappe._dict(tables_json='["lines"]', owner="Administrator", export_format="csv")
			csv_path = Path(directory) / "report.csv"
			csv_count = exports._write(projection, request, csv_path)
			with csv_path.open(encoding="utf-8-sig") as source:
				csv_rows = list(csv.DictReader(source))
			request.export_format = "xlsx"
			xlsx_path = Path(directory) / "report.xlsx"
			xlsx_count = exports._write(projection, request, xlsx_path)
			workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
			rows = list(workbook["库存明细"].values)
			workbook.close()
			self.assertEqual(csv_count, xlsx_count)
			self.assertEqual(len(csv_rows), len(rows) - 1)
			self.assertEqual(float(csv_rows[0]["stock_qty"]), 12)
			self.assertEqual(rows[1][rows[0].index("stock_qty")], 12)
			self.assertEqual(projection.count("conversions"), 1)
			self.assertEqual(projection.count("items"), 1)

	def test_large_native_draft_export_preserves_all_1005_allocations_and_stable_order(self):
		data = self.movement(quantity=1)
		data["items"] = [copy.deepcopy(data["items"][0]) for _ in range(1005)]
		doc = self.create(data)
		query = self.query(
			"movements", "lines", record=doc.name, statuses=[0], sortBy="item_name", sort="asc"
		)
		preview = report_page(query, page_length=1)
		self.assertEqual(preview["matched_count"], 1005)
		self.assertEqual(len(preview["rows"]), 1)
		with tempfile.TemporaryDirectory() as directory, Projection(query) as projection:
			projection.build()
			request = frappe._dict(tables_json='["lines"]', owner="Administrator", export_format="csv")
			csv_path = Path(directory) / "large.csv"
			self.assertEqual(exports._write(projection, request, csv_path), 1005)
			with csv_path.open(encoding="utf-8-sig") as source:
				csv_rows = list(csv.DictReader(source))
			request.export_format = "xlsx"
			xlsx_path = Path(directory) / "large.xlsx"
			self.assertEqual(exports._write(projection, request, xlsx_path), 1005)
			workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
			rows = iter(workbook["库存明细"].values)
			columns = next(rows)
			ids = [row[columns.index("id")] for row in rows]
			workbook.close()
			self.assertEqual(ids, [row["id"] for row in csv_rows])
			self.assertEqual(len(set(ids)), 1005)
			self.assertEqual({row["line_name"] for row in csv_rows}, {row.name for row in doc["items"]})
			self.assertEqual(sum(float(row["stock_qty"]) for row in csv_rows), 1005)

	def test_export_request_is_owned_queued_and_failed_retry_retains_scope(self):
		self.receipt()
		query = self.query()
		with patch("temple_inventory.exports._enqueue") as enqueue:
			result = exports.request_export(query, "xlsx", ["summary", "locations"])
			self.assertEqual(result["status"], "Queued")
			enqueue.assert_called_once()
			doc = frappe.get_doc("Inventory Export Request", result["name"])
			self.assertEqual(doc.owner, "Administrator")
			doc.status = "Failed"
			exports._save(doc)
			retried = exports.retry_export(doc.name, str(doc.modified))
			self.assertEqual(retried["attempt"], 2)
			self.assertEqual(json.loads(doc.query_json), query)
		with self.assertRaises(frappe.ValidationError):
			exports.request_export(query, "csv", ["summary", "locations"])

	def test_export_private_download_rechecks_current_warehouse_and_request_owner(self):
		from frappe.cache_manager import clear_defaults_cache

		self.receipt()
		user = frappe.get_doc(
			{
				"doctype": "User",
				"email": f"ti-export-{self.token}@example.test",
				"first_name": "Export manager",
				"send_welcome_email": 0,
				"user_type": "System User",
				"roles": [{"role": "System Manager"}, {"role": "Stock Manager"}, {"role": "Stock User"}],
			}
		).insert()
		self.addCleanup(clear_defaults_cache, user.name)
		frappe.set_user(user.name)
		query = self.query()
		with patch("temple_inventory.exports._enqueue"):
			result = exports.request_export(query, "csv", ["summary"])
		request = frappe.get_doc("Inventory Export Request", result["name"])
		request.status = "Running"
		exports._save(request)
		with tempfile.TemporaryDirectory() as directory, Projection(query) as projection:
			projection.build()
			output = Path(directory) / "export.csv"
			manifest = Path(directory) / "access.jsonl"
			exports._write(projection, request, output)
			exports._manifest(projection, manifest)
			request.output_file = exports._private_file(output, request, "export.csv")
			request.access_manifest = exports._private_file(manifest, request, "access.jsonl")
		for field in ("output_file", "access_manifest"):
			file = frappe.get_doc("File", request.get(field))
			self.addCleanup(Path(file.get_full_path()).unlink, missing_ok=True)
		request.status = "Completed"
		exports._save(request)
		exports.validate_download(request)
		output_file = frappe.get_doc("File", request.output_file)
		self.assertTrue(output_file.is_downloadable())
		self.assertFalse(frappe.get_doc("File", request.access_manifest).is_downloadable())
		frappe.set_user("Administrator")
		with self.assertRaises(frappe.PermissionError):
			exports.export_status(request.name)
		frappe.get_doc(
			{"doctype": "User Permission", "user": user.name, "allow": "Warehouse", "for_value": self.b}
		).insert()
		frappe.clear_cache(user=user.name)
		frappe.set_user(user.name)
		with self.assertRaises(frappe.PermissionError):
			exports.validate_download(request)
		self.assertFalse(output_file.is_downloadable())

	def test_worker_uses_request_owner_and_fails_if_account_is_disabled_during_generation(self):
		from frappe.cache_manager import clear_defaults_cache

		self.receipt()
		user = frappe.get_doc(
			{
				"doctype": "User",
				"email": f"ti-worker-{self.token}@example.test",
				"first_name": "Worker owner",
				"send_welcome_email": 0,
				"user_type": "System User",
				"roles": [{"role": "System Manager"}, {"role": "Stock Manager"}, {"role": "Stock User"}],
			}
		).insert()
		self.addCleanup(clear_defaults_cache, user.name)
		frappe.set_user(user.name)
		with patch("temple_inventory.exports._enqueue"):
			request = exports.request_export(self.query(), "csv", ["summary"])
		frappe.set_user("Administrator")
		write = exports._write

		def revoke_after_write(projection, doc, path):
			self.assertEqual(frappe.session.user, user.name)
			count = write(projection, doc, path)
			frappe.db.set_value("User", user.name, "enabled", 0)
			return count

		# Exercise the real worker/permission/write pipeline in this test's outer
		# savepoint; suppress only its transaction boundaries, never authorization.
		with (
			patch("frappe.db.commit"),
			patch("frappe.db.rollback"),
			patch("temple_inventory.exports._write", side_effect=revoke_after_write) as writer,
			patch("temple_inventory.exports._private_file") as private_file,
		):
			exports.generate_export(request["name"])
		self.assertEqual(frappe.session.user, "Administrator")
		writer.assert_called_once()
		private_file.assert_not_called()
		failed = frappe.get_doc("Inventory Export Request", request["name"])
		self.assertEqual(failed.status, "Failed")
		self.assertFalse(failed.output_file)
		self.assertFalse(failed.access_manifest)

	def test_worker_discards_cached_permissions_when_owner_role_is_revoked_during_generation(self):
		from frappe.cache_manager import clear_defaults_cache

		self.receipt()
		user = frappe.get_doc(
			{
				"doctype": "User",
				"email": f"ti-role-worker-{self.token}@example.test",
				"first_name": "Role worker owner",
				"send_welcome_email": 0,
				"user_type": "System User",
				"roles": [{"role": "System Manager"}, {"role": "Stock Manager"}, {"role": "Stock User"}],
			}
		).insert()
		self.addCleanup(clear_defaults_cache, user.name)
		frappe.set_user(user.name)
		with patch("temple_inventory.exports._enqueue"):
			request = exports.request_export(self.query(), "csv", ["summary"])
		frappe.set_user("Administrator")
		write = exports._write

		def revoke_after_write(projection, doc, path):
			self.assertIn("System Manager", frappe.get_roles())
			count = write(projection, doc, path)
			local_cache = copy.deepcopy(frappe.local.cache)
			role_permissions = copy.deepcopy(frappe.local.role_permissions)
			# Native role edit invalidates shared Redis caches. Retain the running
			# worker's old local memoization, as a change in another request would.
			frappe.set_user("Administrator")
			changed = frappe.get_doc("User", user.name)
			changed.set("roles", [{"role": "Stock Manager"}, {"role": "Stock User"}])
			changed.save()
			frappe.set_user(user.name)
			frappe.local.cache = local_cache
			frappe.local.role_permissions = role_permissions
			self.assertIn("System Manager", frappe.get_roles())
			return count

		with (
			patch("frappe.db.commit"),
			patch("frappe.db.rollback"),
			patch("temple_inventory.exports._write", side_effect=revoke_after_write) as writer,
			patch("temple_inventory.exports._private_file") as private_file,
		):
			exports.generate_export(request["name"])
		self.assertEqual(frappe.session.user, "Administrator")
		writer.assert_called_once()
		private_file.assert_not_called()
		failed = frappe.get_doc("Inventory Export Request", request["name"])
		self.assertEqual(failed.status, "Failed")
		self.assertFalse(failed.output_file)
		self.assertFalse(failed.access_manifest)

	def test_loan_record_status_uses_complete_original_before_allocation_filters(self):
		self.receipt(10)
		data = self.movement("Loan", 3, self.a, self.loan_leaf, ti_borrower="Whole Loan")
		data["items"].append(dict(data["items"][0], qty=4))
		loan = self.submit(self.create(data))
		self.submit(self.settlement(loan, 3))
		line = loan["items"][0].name
		query = self.query(
			"loans", "outcomes", loan=loan.name, restrictAllocations=True, allocationKeys=[line]
		)
		open_rows = report_page({**query, "loanStatus": "open"})["rows"]
		self.assertEqual(len(open_rows), 1)
		self.assertEqual(open_rows[0]["outstanding"], 0)
		self.assertEqual(open_rows[0]["loan_status"], "Open")
		self.assertEqual(report_page({**query, "loanStatus": "settled"})["matched_count"], 0)
		complete = report_page(self.query("loans", "outcomes", loan=loan.name, loanStatus="open"))
		self.assertEqual(complete["matched_count"], 2)
		self.assertEqual(sum(row["outstanding"] for row in complete["rows"]), 4)

	def test_draft_and_cancelled_loan_records_have_no_current_outstanding_balance(self):
		self.receipt(5)
		draft = self.create(self.movement("Loan", 3, self.a, self.loan_leaf, ti_borrower="Loan Lifecycle"))
		rows = report_page(self.query("loans", "outcomes", loan=draft.name, statuses=[0]))["rows"]
		self.assertEqual(rows[0]["borrowed"], 3)
		self.assertEqual(rows[0]["outstanding"], 0)
		completed = self.submit(draft)
		native_api.cancel_document("Stock Entry", completed.name, completed.version)
		self.assertEqual(native_api.loan_detail(completed.name)["items"][0]["outstanding_qty"], 0)
		query = self.query("loans", "outcomes", loan=completed.name, statuses=[2])
		self.assertEqual(report_page(query)["rows"][0]["outstanding"], 0)
		self.assertEqual(report_page({**query, "loanStatus": "settled"})["matched_count"], 0)

	def test_date_sort_uses_native_posting_time_before_identity_on_the_same_day(self):
		late = self.submit(self.create(self.movement(quantity=1, posting_time="12:00:00.123456")))
		early = self.submit(self.create(self.movement(quantity=1, posting_time="08:00:00.654321")))
		query = self.query("movements", "records", sortBy="date", sort="asc")
		rows = report_page(query)["rows"]
		self.assertEqual([row["document"] for row in rows], [early.name, late.name])
		self.assertTrue(rows[0]["time"].endswith("654321"))

	def test_adjustment_exclusion_is_canonical_in_browse_and_export_queries(self):
		self.receipt(10)
		draft = counts.save_count(str(uuid4()), self.count_data(8, "Reviewed"))
		completed = native_api.submit_document("Stock Reconciliation", draft.name, draft.version)
		query = self.query("movements", "lines", record=completed.name)
		self.assertEqual(report_page(query)["matched_count"], 1)
		self.assertEqual(report_page({**query, "includeAdjustments": False})["matched_count"], 0)
		with self.assertRaises(frappe.ValidationError):
			validate_query({**query, "includeAdjustments": "false"})


def run(test_name=None):
	# Inherited native cases run once in the native runner, not again here.
	names = (
		[test_name]
		if test_name
		else [name for name in NativeReportTests.__dict__ if name.startswith("test_")]
	)
	result = unittest.TextTestRunner(verbosity=2).run(
		unittest.TestSuite(NativeReportTests(name) for name in sorted(names))
	)
	if not result.wasSuccessful():
		raise AssertionError(
			f"Native reports failed: {len(result.failures)} failures, {len(result.errors)} errors"
		)
	return {"tests": result.testsRun, "failures": 0, "errors": 0}
